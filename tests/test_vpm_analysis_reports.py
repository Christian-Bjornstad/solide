"""The Archer/VPM layout with Solide's analysis and evidence boundaries."""
from dataclasses import replace
from contextlib import closing
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import zipfile

from openpyxl import load_workbook as open_workbook
from PIL import Image, ImageDraw
import pytest

from solide.models import Analysis, Session, Variant
from solide.reporting import SOURCE_ORDER, export_analysis


def load_workbook(path):
    return closing(open_workbook(path))


def make_session(*variants,profile='OPA'):
    analysis=Analysis(id='analysis-a',patient='SAMPLE',profile=profile,assay='Oncomine Precision Assay',
                      assembly='GRCh37',metadata={'source metadata':'retained'})
    rows=[replace(variant,analysis_id=analysis.id,patient=analysis.patient) for variant in variants]
    return Session(variants=rows,analyses=[analysis]),rows


def text(book):
    return '\n'.join(str(cell.value) for sheet in book for row in sheet for cell in row if cell.value is not None)


def capture(tmp_path,name='capture.png',size=(1400,2300)):
    path=tmp_path/name
    image=Image.new('RGB',size,'white')
    ImageDraw.Draw(image).rectangle((0,0,size[0]//2,size[1]),fill='#123456')
    image.save(path)
    return path


def result(variant,path=None,source='Franklin',**changes):
    item={'database':source,'status':'found','fingerprint':variant.fingerprint('Other'),
          'clinical_significance':'Database assertion','url':'https://example.org/record',
          'captured_at':'2026-10-15T08:00:00+00:00','raw':{}}
    if path:item['raw']['screenshots']=[{'path':str(path),'label':'Classification evidence'}]
    item.update(changes)
    return item


def test_normalization_conflict_never_embeds_a_green_normalizer_panel(tmp_path):
    session,rows=make_session(Variant(gene='EGFR',selected=True))
    item=result(rows[0],capture(tmp_path),source='Mutalyzer')
    item['status']='needs_review'
    item['raw']['nomenclature_assessment']={'status':'Conflict','identity_confirmed':False,
        'reason':'Different genomic event','suggested_hgvs':''}
    rows[0].evidence['Mutalyzer']=item
    with load_workbook(export_analysis(session,'analysis-a',tmp_path).report) as book:
        assert not any(sheet._images for sheet in book)
        assert 'Different genomic event' in text(book)


def test_default_is_one_workbook_with_vpm_layout_and_source_order(tmp_path):
    session,rows=make_session(Variant(gene='MET',coding='c.1A>T',protein='p.X',selected=True,af_percent=1),
        Variant(gene='EGFR',coding='c.2A>T',protein='p.Y',selected=True,af_percent=90))
    bundle=export_analysis(session,'analysis-a',tmp_path)
    assert bundle.paths==(bundle.report,) and bundle.attachment is None
    with load_workbook(bundle.report) as book:
        assert book.sheetnames==['Overview','MET','EGFR','Quality','Sources','Raw data']
        assert book['Sources'].sheet_state==book['Raw data'].sheet_state=='hidden'
        overview=book['Overview']
        assert {'A1:M2','A6:M6','D3:M4'}<={str(item) for item in overview.merged_cells.ranges}
        assert overview['A1'].font.sz==22 and overview['A1'].fill.fgColor.rgb[-6:]=='163B5C'
        assert [overview.cell(7,col).value for col in range(1,14)]==['Gene','HGVSc','HGVSp','Short evidence','Comment',*SOURCE_ORDER]
        assert [overview.cell(row,1).value for row in (8,9)]==['MET','EGFR']
        assert [overview.column_dimensions[col].width for col in 'ABCDE']==[14,31,25,52,32]
        assert all(overview.column_dimensions[col].width==22 for col in 'FGHIJKLM')
        assert overview.freeze_panes=='A8' and overview.auto_filter.ref=='A7:M9'
        assert overview.row_dimensions[8].height>=108
        assert overview['A8'].fill.fgColor.rgb[-6:]=='EAF3FA'
        assert overview['A9'].fill.fgColor.rgb[-6:]=='FFFFFF'
        variant=book['MET']
        assert 'A1:L2' in {str(item) for item in variant.merged_cells.ranges}
        assert variant['A5'].value=='Compact evidence and source links'
        assert variant.freeze_panes=='A5'
        assert [variant.cell(6+i,1).value for i in range(8)]==list(SOURCE_ORDER)
        assert 'gnomAD AF' not in text(book) and 'WHO driver gene' not in text(book)


def test_vpm_images_keep_full_pixels_and_reserve_height_without_segmentation(tmp_path):
    session,rows=make_session(Variant(gene='EGFR',selected=True))
    path=capture(tmp_path)
    rows[0].evidence['Franklin']=result(rows[0],path)
    original_hash=sha256(path.read_bytes()).hexdigest()
    bundle=export_analysis(session,'analysis-a',tmp_path)
    with load_workbook(bundle.report) as book:
        image=book['EGFR']._images[0]
        assert len(book['EGFR']._images)==1
        assert image.anchor.ext.cx==1120*9525
        assert image.anchor.ext.cy==1840*9525
        first=image.anchor._from.row+1
        assert book['EGFR'].row_dimensions[first].height==28.5
        assert book['EGFR'].row_dimensions[first+48].height==28.5
        caption=book['EGFR'].cell(first-1,1)
        assert caption.value=='Classification evidence' and caption.hyperlink is None
    with zipfile.ZipFile(bundle.report) as archive:
        pngs=[name for name in archive.namelist() if name.startswith('xl/media/')]
        assert len(pngs)==1
        with Image.open(BytesIO(archive.read(pngs[0]))) as embedded,Image.open(path) as source:
            assert embedded.size==source.size and embedded.tobytes()==source.tobytes()
    assert sha256(path.read_bytes()).hexdigest()==original_hash


def test_vpm_keeps_scope_all_qc_raw_and_app_assessment_distinct(tmp_path):
    note=('Full assessment — long note.\n'*100)+'END OF APP ASSESSMENT'
    session,rows=make_session(Variant(gene='EGFR',selected=True,comment=note,classification='Lab assessment',
        report_decision='Include',reviewer='Lab reviewer',reviewed_at='2026-10-15T09:00:00Z',
        transcript='NM_005228.5',coding='c.1A>T',corrected_hgvs='NC_000007.13:g.1A>T',
        raw={'_columns':['Gene','Literal','Empty'],'_values':['EGFR','=1+1','']}),
        Variant(gene='LOW',kind='snp',coverage=4),Variant(gene='RNA',kind='RNAExonTiles',call=''),
        Variant(gene='RNAABSENT',kind='RNAExonVariant',call='ABSENT'))
    rows[0].evidence['ClinVar']=result(rows[0],source='ClinVar',raw={'payload':'X'*40000})
    other=Analysis(id='other',patient='SAMPLE',profile='Ion Torrent')
    session.analyses.append(other)
    session.variants.append(Variant(analysis_id=other.id,patient='SAMPLE',gene='DO-NOT-MIX',selected=True))
    bundle=export_analysis(session,'analysis-a',tmp_path)
    with load_workbook(bundle.report) as book:
        all_values=text(book)
        assert 'DO-NOT-MIX' not in all_values
        assert 'Lab assessment' in all_values and 'Database assertion' in all_values
        assert book['Overview']['B8'].value=='c.1A>T'
        assert 'NC_000007.13:g.1A>T' in all_values and 'NM_005228.5' in all_values
        assert 'END OF APP ASSESSMENT' in all_values and 'Lab reviewer' in all_values
        quality='\n'.join(str(cell.value) for row in book['Quality'] for cell in row if cell.value is not None)
        assert all(word in quality for word in ['LOW','Failed','RNA','Unknown','RNAABSENT','Report'])
        assert book['Raw data']['B2'].value=='=1+1' and book['Raw data']['B2'].data_type=='s'
        assert book['Raw data'].max_row==5
        sources=book['Sources']
        stored=next(row for row in sources if row[7].value==rows[0].id and row[9].value=='JSON response')
        reconstructed=''.join(str(cell.value or '') for cell in stored[10:])
        assert json.loads(reconstructed)==rows[0].evidence['ClinVar']['raw']
        assert not any(cell.data_type=='f' for sheet in book for row in sheet for cell in row)
        for sheet in book:
            for row in sheet:
                for cell in row:
                    if cell.hyperlink and (cell.hyperlink.target or '').startswith('#'):
                        target=cell.hyperlink.target.split('!')[0][1:].strip("'").replace("''", "'")
                        assert target in book.sheetnames and book[target].sheet_state=='visible'


def test_rejected_or_stale_capture_stays_excluded_and_sources_retained(tmp_path):
    session,rows=make_session(Variant(gene='EGFR',selected=True),Variant(gene='MET',selected=True))
    path=capture(tmp_path,size=(300,300))
    rows[0].evidence['Franklin']=result(rows[0],path)
    rows[0].evidence['Franklin']['raw']['identity_verification']={'accepted':False}
    rows[1].evidence['Franklin']=result(rows[1],path,fingerprint='outdated')
    with load_workbook(export_analysis(session,'analysis-a',tmp_path).report) as book:
        assert not any(sheet._images for sheet in book)
        assert 'Identity mismatch' in text(book) and 'Outdated' in text(book)
        assert len([row for row in book['Sources'] if row[2].value=='Franklin'])==2


@pytest.mark.parametrize('status',['found','not_found'])
def test_mtbp_full_report_once_second_and_display_crop_is_lossless(tmp_path,status):
    from solide.evidence import batch_fingerprint
    session,rows=make_session(Variant(gene='EGFR',selected=True),Variant(gene='MET',selected=True))
    full=capture(tmp_path,'full.png',size=(500,800))
    full.with_suffix('.geometry.json').write_text(json.dumps({'content_top':200,'height':800}))
    before=full.read_bytes()
    for variant in rows:
        item=result(variant,full,source='MTBP',status=status,analysis_id=variant.analysis_id,
            mtbp_batch=batch_fingerprint(session,'SAMPLE','analysis-a'),
            url='https://mtbp.org/patients/private/report/private/')
        item['raw']['patient_report_screenshot']=str(full)
        # A full report referenced in screenshots must not repeat on each variant.
        variant.evidence['MTBP']=item
    bundle=export_analysis(session,'analysis-a',tmp_path)
    with load_workbook(bundle.report) as book:
        assert book.sheetnames[:4]==['Overview','MTBP','EGFR','MET']
        assert book['MTBP']['A1'].value=='SAMPLE'
        assert book['MTBP']['A6'].value=='MTBP – combined report'
        assert len(book['MTBP']._images)==1
        assert len([image for sheet in book for image in sheet._images])==1
        assert book['MTBP']._images[0].anchor._from.row==6
        assert 'https://mtbp.org/patients/' not in text(book)
        assert not any(cell.hyperlink and 'mtbp.org/patients/' in (cell.hyperlink.target or '') for sheet in book for row in sheet for cell in row)
    with zipfile.ZipFile(bundle.report) as archive:
        png=next(name for name in archive.namelist() if name.startswith('xl/media/'))
        with Image.open(BytesIO(archive.read(png))) as embedded,Image.open(full) as source:
            cropped=source.crop((0,200,500,800))
            assert embedded.size==cropped.size and embedded.tobytes()==cropped.tobytes()
    assert full.read_bytes()==before


def test_unassigned_and_no_evidence_export(tmp_path):
    session,rows=make_session(Variant(gene='EGFR',selected=True),profile='Unassigned')
    with pytest.raises(ValueError,match='Choose OPA or Ion Torrent'):
        export_analysis(session,'analysis-a',tmp_path)
    session.analyses[0].profile='OPA'
    rows[0].evidence['Franklin']=result(rows[0],capture(tmp_path))
    with load_workbook(export_analysis(session,'analysis-a',tmp_path,include_evidence=False).report) as book:
        assert not any(sheet._images for sheet in book)
        assert 'Franklin' in text(book) and book['Sources'].sheet_state=='hidden'


def test_structured_normalization_is_visible_and_never_becomes_lab_classification(tmp_path):
    session,rows=make_session(Variant(gene='PTEN',selected=True,classification='Laboratory review'))
    raw={'original':'NM_000314.8:c.100_102delinsA',
         'normalization':{'normalized_description':'NM_000314.8:c.100_101del'},
         'nomenclature_assessment':{'status':'Conflict','reason':'Source alleles do not agree.',
             'identity_confirmed':False,'suggested_hgvs':'NM_000314.8:c.100_101del',
             'genomic_hgvs':'NC_000010.10:g.1_2del','source_genomic_query':'10-1-AC-A'}}
    rows[0].evidence['Mutalyzer']=result(rows[0],source='Mutalyzer',status='needs_review',raw=raw)
    with load_workbook(export_analysis(session,'analysis-a',tmp_path).report) as book:
        fields={row[0].value:row[2].value for row in book['PTEN'] if row[0].value}
        assert fields['Normalization status']=='Conflict'
        assert fields['Original name']==raw['original']
        assert fields['Checked name']==raw['nomenclature_assessment']['suggested_hgvs']
        assert fields['Genomic identity']=='Not confirmed'
        assert fields['Lab classification']=='Laboratory review'


def test_duplicate_gene_names_reservations_and_literal_app_notes(tmp_path):
    session,rows=make_session(Variant(gene='MET',protein='NP_000236.2:p.X1Y',selected=True,comment='=literal app note'),
        Variant(gene='MET',protein='NP_000236.2:p.X1Y',selected=True),Variant(gene='Quality',selected=True))
    bundle=export_analysis(session,'analysis-a',tmp_path,include_evidence=False)
    with load_workbook(bundle.report) as book:
        assert book.sheetnames[:4]==['Overview','MET p.X1Y','MET p.X1Y-2','Quality-2']
        fields={row[0].value:row[2] for row in book['MET p.X1Y'] if row[0].value}
        assert fields['Assessment'].value=='=literal app note' and fields['Assessment'].data_type=='s'


def test_exports_version_without_overwriting_and_atomic_failure_has_no_partial_file(tmp_path,monkeypatch):
    session,rows=make_session(Variant(gene='EGFR',selected=True))
    first=export_analysis(session,'analysis-a',tmp_path).report
    original=first.read_bytes()
    second=export_analysis(session,'analysis-a',tmp_path).report
    assert first!=second and first.read_bytes()==original
    from openpyxl.workbook.workbook import Workbook
    def fail_save(*args,**kwargs):raise PermissionError('File locked')
    monkeypatch.setattr(Workbook,'save',fail_save)
    with pytest.raises(PermissionError,match='retry export without rerunning'):
        export_analysis(session,'analysis-a',tmp_path)
    assert set(tmp_path.glob('*.xlsx'))=={first,second}
    assert not list(tmp_path.glob('.solide-*'))
