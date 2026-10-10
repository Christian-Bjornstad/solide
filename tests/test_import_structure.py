import pytest
from openpyxl import Workbook,load_workbook
from solide.importing import load_file
from solide.models import Session
from solide.reporting import export_patient


def test_preamble_empty_column_and_repeated_header_preserved(tmp_path):
    source=tmp_path/'source.tsv'
    source.write_text('##reference=hg19\nMetadata\nType\tGene\t\tCoverage\nSNV\tEGFR\tOriginal value\t400\nType\tGene\t\tCoverage\n',encoding='utf8')
    data=load_file(source);assert len(data.variants)==1
    v=data.variants[0];v.patient='DEMO'
    assert v.source_row==4 and v.raw['_columns']==['Type','Gene','','Coverage']
    report=load_workbook(export_patient(Session(variants=[v]),'DEMO',tmp_path))
    rows=list(report['Raw data'].values)
    assert rows[0][:4]==('Type','Gene','Unnamed column 3','Coverage')
    assert rows[1][:4]==('SNV','EGFR','Original value','400')


def test_ambiguous_reviewed_sheets_rejected(tmp_path):
    w=Workbook()
    for sheet in [w.active,w.create_sheet('Presentation copy')]:
        sheet.append(['Type','Gene']);sheet.append(['SNV','EGFR'])
    path=tmp_path/'duplicate.xlsx';w.save(path)
    with pytest.raises(ValueError,match='Multiple variant tables'):load_file(path)


def test_duplicate_headers_rejected_instead_of_overwritten(tmp_path):
    path=tmp_path/'duplicate.tsv';path.write_text('Type\tGene\tCoding\tCoding\nSNV\tEGFR\tc.1A>T\tc.2A>G')
    with pytest.raises(ValueError,match='Duplicate column names'):load_file(path)


def test_placeholder_and_provenance_header_names_do_not_overwrite_source(tmp_path):
    source=tmp_path/'columns.tsv'
    source.write_text('Type\tGene\t\tUnnamed column 3\tSource row\nSNV\tEGFR\tBlank-header value\tNamed-header value\tOriginal row value')
    v=load_file(source).variants[0];v.patient='DEMO'
    w=load_workbook(export_patient(Session(variants=[v]),'DEMO',tmp_path))
    rows=list(w['Raw data'].values)
    assert rows[0][:5]==('Type','Gene','Unnamed column 3 (2)','Unnamed column 3','Source row')
    assert rows[1][:5]==('SNV','EGFR','Blank-header value','Named-header value','Original row value')
    assert len(rows[0])==len(set(str(value).casefold() for value in rows[0]))


def test_explicit_solide_report_marker_rejects_one_raw_table(tmp_path):
    w = Workbook()
    w.properties.identifier = 'solide:variant-review-report'
    w.active.title = 'Results'
    w.active.append(['Variant report'])
    raw = w.create_sheet('Raw data')
    raw.append(['Type', 'Gene'])
    raw.append(['SNV', 'EGFR'])
    path = tmp_path / 'review-report.xlsx'
    w.save(path)

    with pytest.raises(ValueError, match='Solide report.*original'):
        load_file(path)


def test_legacy_solide_report_rejects_one_raw_table(tmp_path):
    w = Workbook()
    overview = w.active
    overview.title = 'Overview'
    overview.append(['SOLIDE'])
    overview.append(['Patient / sample', 'SYNTHETIC'])
    overview.append(['Generated (UTC)', None, 'App version', '0.4.2', 'Report format', 2])
    raw = w.create_sheet('Raw data')
    raw.append(['Type', 'Gene'])
    raw.append(['SNV', 'EGFR'])
    path = tmp_path / 'legacy-report.xlsx'
    w.save(path)

    with pytest.raises(ValueError, match='Solide report.*original'):
        load_file(path)


def test_solide_title_alone_does_not_reject_an_original_export(tmp_path):
    w = Workbook()
    sheet = w.active
    sheet.append(['SOLIDE'])
    sheet.append(['Type', 'Gene'])
    sheet.append(['SNV', 'EGFR'])
    path = tmp_path / 'original.xlsx'
    w.save(path)

    assert len(load_file(path).variants) == 1


def test_filtered_ion_scope_warning_survives_a_session_roundtrip(tmp_path):
    from solide.session import load_session, save_session

    path = tmp_path / 'filtered.tsv'
    path.write_text(
        '##reference=hg19\n##totalVariantCount=12\n##filterInCount=2\n'
        '##filteredOutCount=10\n##hiddenVariantCount=0\n'
        'Transcript\tGenes\tCoding\tType\n'
        'NM_000546.6\tTP53\tc.742C>T\tSNV\n'
        'NM_005228.5\tEGFR\tc.2573T>G\tSNV\n', encoding='utf8')

    result = load_file(path)
    warning = next(value for value in result.warnings if value.startswith('Filtered export:'))
    assert '2 imported rows' in warning and '12 total' in warning and '10 filtered out' in warning
    assert 'QC covers imported rows only' in warning
    expected = {
        'total_variant_count': 12, 'exported_variant_count': 2, 'filtered_in_count': 2,
        'filtered_out_count': 10, 'hidden_variant_count': 0, 'filtered_export': True,
    }
    assert all(v.raw['_export_scope'] == expected for v in result.variants)
    assert result.metadata['totalVariantCount'] == '12'
    session_path = tmp_path / 'review.solide.json'
    save_session(Session(variants=result.variants), session_path)
    assert all(v.raw['_export_scope'] == expected for v in load_session(session_path).variants)


@pytest.mark.parametrize('metadata, expected_text', [
    ('##totalVariantCount=12\n', '12 total'),
    ('##filteredOutCount=10\n', '10 filtered out'),
    ('##hiddenVariantCount=1\n', '1 hidden'),
])
def test_partial_scope_metadata_still_exposes_filtered_rows(tmp_path, metadata, expected_text):
    path = tmp_path / 'partial.tsv'
    path.write_text(metadata + 'Type\tGene\nSNV\tEGFR\n', encoding='utf8')
    result = load_file(path)

    warning = next(value for value in result.warnings if value.startswith('Filtered export:'))
    assert expected_text in warning
    assert result.variants[0].raw['_export_scope']['filtered_export'] is True


def test_unfiltered_scope_retains_counts_without_claiming_missing_rows(tmp_path):
    path = tmp_path / 'complete.tsv'
    path.write_text(
        '##totalVariantCount=1\n##filterInCount=1\n##filteredOutCount=0\n'
        '##hiddenVariantCount=0\nType\tGene\nSNV\tEGFR\n', encoding='utf8')
    result = load_file(path)

    assert not any(value.startswith('Filtered export:') for value in result.warnings)
    assert result.variants[0].raw['_export_scope']['filtered_export'] is False
    assert result.variants[0].raw['_export_scope']['total_variant_count'] == 1


@pytest.mark.parametrize('count', ['-1', 'NaN', 'inf', '12.5', 'unknown'])
def test_invalid_scope_counts_are_not_invented(tmp_path, count):
    path = tmp_path / 'malformed.tsv'
    path.write_text(f'##totalVariantCount={count}\nType\tGene\nSNV\tEGFR\n', encoding='utf8')
    result = load_file(path)

    assert not any(value.startswith('Filtered export:') for value in result.warnings)
    assert '_export_scope' not in result.variants[0].raw
    assert result.metadata['totalVariantCount'] == count


def test_source_without_scope_metadata_does_not_claim_complete_export(tmp_path):
    path = tmp_path / 'unknown-scope.tsv'
    path.write_text('Type\tGene\nSNV\tEGFR\n', encoding='utf8')
    result = load_file(path)

    assert '_export_scope' not in result.variants[0].raw
    assert not any(value.startswith('Filtered export:') for value in result.warnings)
