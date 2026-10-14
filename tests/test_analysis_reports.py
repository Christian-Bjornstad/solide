from dataclasses import replace
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import zipfile

from openpyxl import load_workbook
from PIL import Image, ImageDraw
import pytest

from solide.models import Analysis, Session, Variant
from solide.reporting import export_analysis, export_patient


def analysis_session(*variants, profile='OPA', analysis_id='opa'):
    analysis=Analysis(id=analysis_id,patient='SAME-SAMPLE',profile=profile,
                      assay='Oncomine Precision Assay' if profile=='OPA' else 'Childhood',assembly='GRCh37')
    rows=[replace(variant,patient=analysis.patient,analysis_id=analysis_id) for variant in variants]
    return Session(variants=rows,analyses=[analysis]),rows


def all_text(book):
    return '\n'.join(str(cell.value) for sheet in book for row in sheet for cell in row if cell.value is not None)


def source_capture(tmp_path,name='capture.png',height=220):
    image=Image.new('RGB',(400,height),'white')
    ImageDraw.Draw(image).rectangle((0,0,150,height),fill='#425B3D')
    path=tmp_path/name;image.save(path);return path


def evidence(variant,path,source='Franklin',status='found'):
    return {'database':source,'status':status,'fingerprint':variant.fingerprint('Other'),
            'clinical_significance':'Database assertion','captured_at':'2026-10-14T10:00:00+00:00',
            'url':'https://mtbp.org/patients/123/report/456/' if source=='MTBP' else 'https://example.org/record',
            'raw':{'screenshots':[{'path':str(path),'label':'Classification overview'}]}}


def test_compact_analysis_report_and_complete_literal_sources(tmp_path):
    original={'Gene':'EGFR','Type':'snp','Call':'PRESENT','Literal':'=1+1','Empty':'',
              '_columns':['Gene','Type','Call','Literal','Empty'],
              '_values':['EGFR','snp','PRESENT','=1+1','']}
    session,rows=analysis_session(Variant(gene='EGFR',transcript='NM_005228.5',coding='c.2573T>G',
        protein='p.L858R',kind='snp',selected=True,af_percent=12.5,coverage=1000,
        classification='Laboratory assessment',report_decision='Include',raw=original,source_row=3,source_hash='abc'))
    rows[0].evidence['ClinVar']={'database':'ClinVar','status':'needs_review','fingerprint':rows[0].fingerprint('Other'),
                                'clinical_significance':'Different database assertion'}
    bundle=export_analysis(session,'opa',tmp_path,include_evidence=False)
    assert bundle.paths==(bundle.report,) and bundle.attachment is None
    book=load_workbook(bundle.report)
    assert [sheet.title for sheet in book if sheet.sheet_state=='visible']==['Findings','Quality','Sources']
    assert book['Raw data'].sheet_state=='hidden' and book['Findings'].max_column<=8
    assert book.properties.identifier=='solide:variant-review-report'
    assert 'Unhide' in all_text(book)
    assert book['Raw data']['D2'].value=='=1+1' and book['Raw data']['D2'].data_type=='s'
    assert list(book['Raw data'].values)[0][:5]==('Gene','Type','Call','Literal','Empty')
    assert any(cell.value=='Laboratory assessment' for row in book['Findings'] for cell in row)
    assert any(cell.value=='Different database assertion' for row in book['Sources'] for cell in row)
    assert not any(sheet._images for sheet in book)
    assert not any(cell.data_type=='f' for sheet in book for row in sheet for cell in row)
    for sheet in book:
        for row in sheet:
            for cell in row:
                assert not (cell.hyperlink and (cell.hyperlink.target or '').startswith("#'Raw data'"))
    finding_row=next(cell.row for row in book['Findings'] for cell in row if cell.value=='EGFR')
    assert book['Findings'].row_dimensions[finding_row].height<=44
    book.close()


def test_same_sample_analysis_scope_never_combines_genes_or_images(tmp_path):
    session,rows=analysis_session(Variant(gene='EGFR',coding='c.1A>T',selected=True))
    other=Analysis(id='ion',patient='SAME-SAMPLE',profile='Ion Torrent',assay='Childhood',assembly='GRCh37')
    secret=Variant(patient='SAME-SAMPLE',analysis_id='ion',gene='OTHER-ANALYSIS-GENE',selected=True)
    session.analyses.append(other);session.variants.append(secret)
    rows[0].evidence['Franklin']=evidence(rows[0],source_capture(tmp_path))
    secret.evidence['Franklin']=evidence(secret,source_capture(tmp_path,'other.png'))
    bundle=export_analysis(session,'opa',tmp_path)
    assert bundle.attachment and bundle.paths==(bundle.report,bundle.attachment)
    for path in bundle.paths:
        book=load_workbook(path)
        assert 'OTHER-ANALYSIS-GENE' not in all_text(book)
        book.close()
    with pytest.raises(ValueError,match='[Cc]hoose.*analysis'):
        export_patient(session,'SAME-SAMPLE',tmp_path)


def test_quality_keeps_unselected_opa_fail_review_report_and_unknown(tmp_path):
    session,rows=analysis_session(Variant(gene='FINDING',selected=True),
        Variant(gene='LOW',kind='snp',coverage=76,source_row=20),
        Variant(gene='CNVFAIL',kind='CNV',copy_number=.9),
        Variant(gene='CNVBOUNDARY',kind='CNV',copy_number=1),
        Variant(gene='CNVMISSING',kind='CNV'),
        Variant(gene='IMBALANCE',kind='RNAExonTiles',call='NO CALL'),
        Variant(gene='RNAABSENT',kind='RNAExonVariant',call='ABSENT'))
    bundle=export_analysis(session,'opa',tmp_path,include_evidence=False)
    book=load_workbook(bundle.report);text=all_text(book)
    assert all(gene in text for gene in ['LOW','CNVFAIL','CNVBOUNDARY','CNVMISSING','IMBALANCE','RNAABSENT'])
    quality='\n'.join(str(cell.value) for row in book['Quality'] for cell in row)
    assert all(status in quality for status in ['Failed','Review','Report','Unknown'])
    assert 'No exceptions in supplied rows' not in quality
    book.close()


def test_ion_qc_scope_and_identity_checks_do_not_claim_whole_gene_success(tmp_path):
    session,rows=analysis_session(Variant(gene='MET',transcript='NM_001127500.3',coding='c.5346+5C>A',selected=True,
        coverage=2000,raw={'_export_scope':{'filtered_export':True,'total_variant_count':3385,'exported_variant_count':1}}),
        profile='Ion Torrent',analysis_id='ion')
    book=load_workbook(export_analysis(session,'ion',tmp_path,include_evidence=False).report)
    text=all_text(book)
    assert '3,385' in text and 'supplied rows only' in text
    assert 'SpliceAI' in text and 'MANE' in text
    assert 'Not supplied' in text and 'whole-gene' in text
    book.close()


def test_attachment_omits_stale_and_identity_mismatch_images(tmp_path):
    session,rows=analysis_session(Variant(gene='EGFR',coding='c.1A>T',selected=True),
                                Variant(gene='TP53',coding='c.2G>A',selected=True))
    path=source_capture(tmp_path)
    rows[0].evidence['Franklin']=evidence(rows[0],path,status='identity_mismatch')
    rows[1].evidence['ClinVar']={**evidence(rows[1],path,source='ClinVar'),'fingerprint':'old'}
    bundle=export_analysis(session,'opa',tmp_path)
    assert bundle.attachment
    for path in bundle.paths:
        with zipfile.ZipFile(path) as archive:
            assert not any(name.startswith('xl/media/') for name in archive.namelist())
    main=load_workbook(bundle.report)
    assert 'Identity mismatch' in all_text(main) and 'Outdated' in all_text(main)
    main.close()


def test_full_mtbp_report_deduplicated_trimmed_without_changing_original(tmp_path):
    from solide.evidence import batch_fingerprint
    session,rows=analysis_session(Variant(gene='EGFR',coding='c.1A>T',selected=True),
                                Variant(gene='TP53',coding='c.2G>A',selected=True))
    path=source_capture(tmp_path,height=900)
    path.with_suffix('.geometry.json').write_text(json.dumps({'height':900,'content_top':240}),encoding='utf8')
    before=sha256(path.read_bytes()).hexdigest()
    for variant in rows:
        item=evidence(variant,path,source='MTBP');item['mtbp_batch']=batch_fingerprint(session,'SAME-SAMPLE')
        item['analysis_id']=variant.analysis_id
        item['raw']['patient_report_screenshot']=str(path)
        variant.evidence['MTBP']=item
    bundle=export_analysis(session,'opa',tmp_path)
    assert sha256(path.read_bytes()).hexdigest()==before
    attachment=load_workbook(bundle.attachment)
    assert 'MTBP' in attachment.sheetnames and len(attachment['MTBP']._images)==1
    image=attachment['MTBP']._images[0]
    with Image.open(BytesIO(image._data())) as embedded:
        with Image.open(path) as original:
            assert embedded.size==(400,660)
            assert embedded.tobytes()==original.crop((0,240,400,900)).tobytes()
    assert 'https://mtbp.org/patients/' not in all_text(attachment)
    main=load_workbook(bundle.report)
    assert 'https://mtbp.org/patients/' not in all_text(main)
    assert any(cell.hyperlink and cell.hyperlink.target.startswith(bundle.attachment.name)
               for row in main['Sources'] for cell in row)
    main.close();attachment.close()


def test_bad_mtbp_geometry_keeps_entire_source_image(tmp_path):
    from solide.report_layout import mtbp_content_top
    path=source_capture(tmp_path)
    for bad in [{'height':220,'content_top':300},{'height':220,'content_top':float('nan')},
                {'height':0,'content_top':100},{'height':220,'content_top':-5}]:
        path.with_suffix('.geometry.json').write_text(json.dumps(bad),encoding='utf8')
        assert mtbp_content_top(path)==0
    path.with_suffix('.geometry.json').write_text(json.dumps({'height':220,'rows':[{'section':{'y':50}}]}),encoding='utf8')
    assert mtbp_content_top(path)==50


def test_failed_bundle_publication_preserves_existing_files_and_leaves_no_orphans(tmp_path,monkeypatch):
    import solide.reporting as reporting
    session,rows=analysis_session(Variant(gene='EGFR',coding='c.1A>T',selected=True))
    rows[0].evidence['Franklin']=evidence(rows[0],source_capture(tmp_path))
    previous=tmp_path/'existing.xlsx';previous.write_bytes(b'existing report')
    original=reporting.os.replace
    def fail_main(source,target):
        if 'Evidence' not in Path(target).name:raise PermissionError('locked')
        return original(source,target)
    monkeypatch.setattr(reporting.os,'replace',fail_main)
    with pytest.raises(PermissionError,match='retry.*export'):
        export_analysis(session,'opa',tmp_path)
    assert previous.read_bytes()==b'existing report'
    assert [path.name for path in tmp_path.glob('*.xlsx')]==['existing.xlsx']
    assert not list(tmp_path.glob('.solide-*'))


@pytest.mark.parametrize('kind,call',[('RNAExonTiles',''),('RNAExonTiles','MAYBE'),
                                     ('RNAExonVariant',''),('RNAExonVariant','NO CALL')])
def test_supplied_rna_missing_or_unknown_call_is_never_clear(tmp_path,kind,call):
    from solide.quality import qc_flags
    session,rows=analysis_session(Variant(gene='RNA',kind=kind,call=call))
    assert any(flag.status=='Unknown' for flag in qc_flags(rows[0]))
    book=load_workbook(export_analysis(session,'opa',tmp_path,include_evidence=False).report)
    quality='\n'.join(str(cell.value) for row in book['Quality'] for cell in row)
    assert 'Unknown' in quality and 'missing or unrecognized' in quality
    book.close()


def test_raw_identity_rejection_overrides_found_status_for_all_capture_links(tmp_path):
    session,rows=analysis_session(Variant(gene='EGFR',selected=True))
    item=evidence(rows[0],source_capture(tmp_path));item['raw']['identity_verification']={'accepted':False}
    rows[0].evidence['Franklin']=item
    for include in [False,True]:
        bundle=export_analysis(session,'opa',tmp_path,include_evidence=include)
        for path in bundle.paths:
            book=load_workbook(path)
            assert not any(sheet._images for sheet in book)
            assert 'Identity mismatch' in all_text(book)
            if not include:
                assert not any(cell.value=='Open local capture' for row in book['Sources'] for cell in row)
            book.close()


def test_full_source_json_and_long_assertion_are_lossless_literal_chunks(tmp_path):
    session,rows=analysis_session(Variant(gene='BRCA1',selected=True))
    value='X'*32753+'=+'+'Y'*20000
    original={'response':value,'release':{'id':76}}
    assertion='Detailed source assertion. '*1800
    rows[0].evidence['BRCA Exchange']={'database':'BRCA Exchange','status':'needs_review',
        'fingerprint':rows[0].fingerprint('Other'),'raw':original,'clinical_significance':assertion}
    book=load_workbook(export_analysis(session,'opa',tmp_path,include_evidence=False).report)
    source=book['Sources']
    storage=[row for row in source if len(row)>10 and row[9].value=='JSON response']
    assert len(storage)==1
    chunks=[cell.value or '' for cell in storage[0][10:]]
    assert json.loads(''.join(chunks))==original
    assert all(cell.data_type=='s' for cell in storage[0][10:] if cell.value is not None)
    assertion_rows=[row for row in source if len(row)>10 and row[9].value=='Full source assertion']
    assert ''.join(str(cell.value or '') for cell in assertion_rows[0][10:])==assertion
    assert ''.join(str(row[1].value or '') for row in source if row[0].value in
                   {'Full source assertion','Full source assertion (continued)'})==assertion
    assert source.row_dimensions[storage[0][0].row].hidden
    book.close()


def test_unassigned_profile_blocks_export_without_writing_files(tmp_path):
    session,rows=analysis_session(Variant(gene='EGFR',selected=True),profile='Unassigned')
    with pytest.raises(ValueError,match='Choose OPA or Ion Torrent'):
        export_analysis(session,'opa',tmp_path)
    assert not list(tmp_path.glob('*.xlsx'))


def test_reviewed_genomic_variant_has_explicit_source_transcript_label(tmp_path):
    session,rows=analysis_session(Variant(gene='EGFR',transcript='NM_005228.5',coding='c.2573T>G',
        corrected_hgvs='NC_000007.13:g.55259515T>G',selected=True))
    book=load_workbook(export_analysis(session,'opa',tmp_path,include_evidence=False).report)
    assert book['Findings']['B5'].value=='Source transcript'
    assert book['Findings']['C5'].value=='Reviewed / source variant'
    assert book['Findings']['B6'].value=='NM_005228.5'
    assert book['Findings']['C6'].value=='NC_000007.13:g.55259515T>G'
    book.close()


def test_quality_identity_section_uses_the_same_wide_reason_column(tmp_path):
    session,rows=analysis_session(Variant(gene='ARID1B,LOC115308161',transcript='NM_001371656.1,NR_163974.1',
        coding='c.903_904delCG',selected=True,kind='INDEL',coverage=76),profile='Ion Torrent',analysis_id='ion')
    book=load_workbook(export_analysis(session,'ion',tmp_path,include_evidence=False).report)
    sheet=book['Quality']
    header=next(cell.row for row in sheet for cell in row if cell.value=='Required checks')
    assert sheet.cell(header,5).value=='Required checks'
    assert sheet.cell(header+1,5).value=='Multiple transcripts / genes'
    assert sheet.column_dimensions['E'].width==43
    assert sheet.cell(header+1,6).value==rows[0].source_row
    assert sheet.row_dimensions[header+1].height>=44
    book.close()


def test_findings_metadata_height_fits_opa_assay_name(tmp_path):
    session,rows=analysis_session(Variant(gene='EGFR',selected=True))
    book=load_workbook(export_analysis(session,'opa',tmp_path,include_evidence=False).report)
    sheet=book['Findings']
    assert sheet['F2'].value=='Oncomine Precision Assay'
    assert sheet.row_dimensions[2].height>=58
    book.close()


def test_analysis_metadata_is_complete_hidden_literal_storage(tmp_path):
    session,rows=analysis_session(Variant(gene='EGFR',selected=True))
    metadata={'analysisName':'Private machine analysis identifier','workflowName':'OPA',
              'response':'X'*40000,'analysisDate':'2026-10-14'}
    session.analyses[0].metadata=metadata
    book=load_workbook(export_analysis(session,'opa',tmp_path,include_evidence=False).report)
    sheet=book['Sources']
    storage=[row for row in sheet if len(row)>10 and row[9].value=='Source metadata']
    assert len(storage)==1
    assert storage[0][7].value=='opa' and storage[0][8].value=='analysis'
    assert sheet.row_dimensions[storage[0][0].row].hidden
    assert json.loads(''.join(str(cell.value or '') for cell in storage[0][10:]))==metadata
    assert not any(cell.value=='Source metadata' for row in sheet for cell in row[:7])
    assert not any('Private machine analysis identifier' in str(cell.value) for row in sheet
                   if not sheet.row_dimensions[row[0].row].hidden for cell in row[:7])
    assert 'Report format' not in all_text(book)
    book.close()


def test_brca_release_attachment_is_compact_but_source_response_remains_complete(tmp_path):
    session,rows=analysis_session(Variant(gene='BRCA1',coding='c.1054G>T',selected=True))
    release={'id':76,'version':31,'date':'2026-03-08','archive':'release-03-08-26.tar.gz',
             'notes':'Long technical notes and contact@example.org that should not fill the report.'}
    raw={'release':release,'classification':{'expert':'Pathogenic'}}
    rows[0].evidence['BRCA Exchange']={'database':'BRCA Exchange','status':'needs_review',
        'fingerprint':rows[0].fingerprint('Other'),'raw':raw}
    bundle=export_analysis(session,'opa',tmp_path)
    attachment=load_workbook(bundle.attachment)
    text=all_text(attachment)
    assert 'Release: 76' in text and 'Version: 31' in text and 'Date: 2026-03-08' in text
    assert 'Archive: release-03-08-26.tar.gz' in text
    assert 'contact@example.org' not in text and 'Long technical notes' not in text
    main=load_workbook(bundle.report)
    stored=next(row for row in main['Sources'] if len(row)>10 and row[9].value=='JSON response')
    assert json.loads(''.join(str(cell.value or '') for cell in stored[10:]))==raw
    attachment.close();main.close()
