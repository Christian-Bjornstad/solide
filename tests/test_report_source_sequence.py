from contextlib import closing

from openpyxl import load_workbook

from solide.models import Variant
from solide.reporting import export_analysis
from test_vpm_analysis_reports import make_session, capture, result


ORDER=('Mutalyzer','SpliceAI','BRCA Exchange','OncoKB','COSMIC','ClinVar','MTBP','Franklin')


def full_results(tmp_path):
    session,rows=make_session(Variant(gene='BRCA1',transcript='NM_007294.4',coding='c.1+1A>G',
        locus='chr17:41276046',ref='T',alt='C',assembly='GRCh37',selected=True))
    v=rows[0]; image=capture(tmp_path,size=(400,300))
    for source in ORDER:v.evidence[source]=result(v,image,source=source)
    v.evidence['Mutalyzer']['status']='needs_review'
    v.evidence['Mutalyzer']['raw']['nomenclature_assessment']={'status':'Normalized','identity_confirmed':True}
    v.evidence['BRCA Exchange']['raw'].update(assembly_verified='GRCh37',matched_location={'position':41276046})
    splice=v.evidence['SpliceAI']['raw']; query='chr17-41276046-T-C'
    splice.update(query=query,response={'variant':query,'chrom':'17','pos':41276046,'ref':'T','alt':'C',
        'hg':37,'genomeVersion':'37','distance':500,'mask':1,'scores':[{
            **dict.fromkeys(('DS_AG','DS_AL','DS_DG','DS_DL'),'0.1'),
            **dict.fromkeys(('DP_AG','DP_AL','DP_DG','DP_DL'),0)}]},
        screenshots=[{'path':str(image),'label':'Query'},{'path':str(image),'label':'Scores'}],
        capture={'query':query,'score_identity_confirmed':True,'dom_identity_confirmed':True,
            'capture_version':2,'mane_transcripts':['ENST1.2'],
            'panels':{key:{'visible':True,'transcripts':['ENST1.2']} for key in ('spliceai','pangolin','consequences')},
            'settings':{'hg':37,'distance':500,'mask':1,'bc':'basic','ref_alt':True}})
    return session,v


def section_names(sheet):
    return [str(row[0].value).split(' – ')[0] for row in sheet
        if str(row[0].value).split(' – ')[0] in ORDER]


def test_report_columns_and_variant_sections_follow_requested_order(tmp_path):
    session,v=full_results(tmp_path)
    with closing(load_workbook(export_analysis(session,'analysis-a',tmp_path).report)) as book:
        assert tuple(cell.value for cell in book['Oversikt'][7][5:])==ORDER
        assert section_names(book['BRCA1'])==list(ORDER)


def test_normalizer_unchanged_and_unmatched_optional_sources_stay_only_in_audit(tmp_path):
    session,v=full_results(tmp_path)
    v.evidence['Mutalyzer']['raw']['nomenclature_assessment']['status']='Unchanged'
    v.evidence['SpliceAI']['status']='not_found'
    v.evidence['BRCA Exchange']['status']='not_found'
    with closing(load_workbook(export_analysis(session,'analysis-a',tmp_path).report)) as book:
        assert section_names(book['BRCA1'])==list(ORDER[3:])
        assert {row[2].value for row in book['Kilder']}>=set(ORDER)


def test_report_uses_confirmed_corrected_hgvs_and_keeps_raw_source(tmp_path):
    session,rows=make_session(Variant(gene='EGFR',coding='c.2235_2250delinsA',
        transcript='NM_005228.5',corrected_hgvs='NM_005228.5:c.2235_2249del',
        nomenclature_verified=True,selected=True,raw={'Coding':'c.2235_2250delinsA'}))
    with closing(load_workbook(export_analysis(session,'analysis-a',tmp_path).report)) as book:
        assert book['Oversikt']['B8'].value=='NM_005228.5:c.2235_2249del'
        assert 'NM_005228.5:c.2235_2249del' in book['EGFR']['A1'].value
        assert book['Rådata']['A2'].value=='c.2235_2250delinsA'


def test_rechecked_corrected_name_stays_normalized_relative_to_import(tmp_path):
    session,v=full_results(tmp_path)
    v.corrected_hgvs='NM_007294.4:c.2+1A>G';v.nomenclature_verified=True
    item=v.evidence['Mutalyzer'];item['fingerprint']=v.fingerprint('Other')
    item['raw'].update(original=v.corrected_hgvs)
    item['raw']['nomenclature_assessment'].update(status='Unchanged',suggested_hgvs=v.corrected_hgvs)
    with closing(load_workbook(export_analysis(session,'analysis-a',tmp_path).report)) as book:
        assert book['Oversikt']['F8'].value=='Normalisert'
        assert section_names(book['BRCA1'])[0]=='Mutalyzer'
