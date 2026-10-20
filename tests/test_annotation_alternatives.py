"""One genomic source row can expose distinct, selectable transcript annotations."""
from copy import deepcopy
from contextlib import closing
import pytest
from solide.models import Variant,Session,ImportResult
from solide.analyses import add_analysis
from solide.evidence import query_record,build_search_plan
from solide.reporting import export_analysis
from openpyxl import load_workbook
from solide.annotation_resolution import resolve_annotation
from solide.annotations import annotation_resolution_current
from solide.nomenclature import hgvs_query,genomic_query

def imported(**changes):
    fields=dict(gene='KLLN,PTEN',transcript='NM_001126049.2, NM_000314.8',
        coding='c.-2001C>T, c.19G>A',protein='p.?, p.Glu7Lys',assembly='GRCh37',
        locus='chr10:89624245',ref='G',alt='A',kind='SNV',call='PRESENT',
        source_file='synthetic.tsv',source_hash='synthetic-hash',source_row=20)
    fields.update(changes);v=Variant(**fields)
    columns=['Genes','Transcript','Coding','Amino Acid Change']
    values=[v.gene,v.transcript,v.coding,v.protein]
    v.raw=dict(zip(columns,values));v.raw.update(_columns=columns,_values=values)
    return v

def session_with(v):
    s=Session();a=add_analysis(s,ImportResult([v]),'00OUM00001','GRCh37','Ion Torrent','Childhood')
    return s,a


def confirm(v):
    """Inject an independent reference/event result; never issue live test requests."""
    expected=hgvs_query(v);gene=v.gene
    class References:
        def resolve(self,refs):return {r:{'status':'resolved','transcript':r,'gene':gene} for r in refs}
    def normalize(probe):
        full=hgvs_query(probe)
        return {'original':full,'normalization':{'input_description':full,'normalized_description':full,
                    'gene_id':gene,'normalized_model':{}},
            'nomenclature_assessment':{'identity_confirmed':full==expected,
                'status':'Unchanged' if full==expected else 'Conflict','source_genomic_query':genomic_query(probe)},
            'genomic_identity_check':{'matched_event':{'reference':'NC_000010.10','start':89624244,'end':89624245,'inserted':'A'}}}
    assert resolve_annotation(v,References(),normalize=normalize)
    assert annotation_resolution_current(v)

def test_aligned_annotations_split_without_cross_product_or_shared_state():
    original=imported();raw=deepcopy(original.raw);s,a=session_with(original)
    assert [(v.gene,v.transcript,v.coding,v.protein) for v in s.variants]==[
        ('KLLN','NM_001126049.2','c.-2001C>T','p.?'),
        ('PTEN','NM_000314.8','c.19G>A','p.Glu7Lys')]
    assert len({v.id for v in s.variants})==2
    assert not any(v.selected for v in s.variants)
    for v in s.variants:
        assert v.source_row==20 and v.analysis_id==a.id
        assert v.raw['_columns']==raw['_columns'] and v.raw['_values']==raw['_values']
    assert original.raw==raw
    s.variants[1].selected=True;s.variants[1].evidence['ClinVar']={'status':'not_found'}
    assert not s.variants[0].selected and not s.variants[0].evidence
    assert build_search_plan(s,['ClinVar'],'all')=={(s.variants[1].id,'ClinVar')}
    with pytest.raises(ValueError):query_record(s.variants[1],'ClinVar')
    confirm(s.variants[1]);record=query_record(s.variants[1],'ClinVar')
    assert record.symbol=='PTEN' and record.hgvsc=='NM_000314.8:c.19G>A'

@pytest.mark.parametrize('gene,transcript,coding,protein',[
    ('TET2,TET2-AS1','NM_001127208.3, NR_126420.1','c.2815C>T','p.Gln939Ter'),
    ('ARID1B,LOC115308161','NM_001371656.1, NR_163974.1','c.903_904delCG','p.Gly303ArgfsTer11'),
])
def test_single_coding_annotation_is_never_copied_to_noncoding_transcript(gene,transcript,coding,protein):
    s,a=session_with(imported(gene=gene,transcript=transcript,coding=coding,protein=protein))
    first,second=s.variants
    assert first.coding==coding and first.protein==protein
    assert second.transcript.startswith('NR_') and second.coding==second.protein==''
    with pytest.raises(ValueError):query_record(second,'ClinVar')

def test_nonparallel_ambiguous_annotations_are_not_guessed():
    s,a=session_with(imported(transcript='NM_1.1, NM_2.2',coding='c.1A>T',protein='p.Lys1Asn'))
    assert len(s.variants)==2
    assert all(v.coding=='' for v in s.variants)
    for v in s.variants:
        with pytest.raises(ValueError):query_record(v,'ClinVar')

def test_selected_alternative_exports_one_finding_and_one_original_raw_row(tmp_path):
    source=imported();s,a=session_with(source);s.variants[1].selected=True
    with pytest.raises(ValueError):export_analysis(s,a.id,tmp_path)
    confirm(s.variants[1])
    with closing(load_workbook(export_analysis(s,a.id,tmp_path).report)) as book:
        assert book['Oversikt'].max_row==8
        assert book['Oversikt']['A8'].value=='PTEN'
        assert book['Oversikt']['B8'].value=='c.19G>A'
        assert book['Rådata'].max_row==2
        assert [cell.value for cell in book['Rådata'][2]][:4]==source.raw['_values']
    # Repeat import must reuse the existing choices and the two child IDs.
    ids=[v.id for v in s.variants]
    assert add_analysis(s,ImportResult([source]),a.patient,a.assembly,a.profile,a.assay) is a
    assert [v.id for v in s.variants]==ids and s.variants[1].selected


def test_reversed_noncoding_reference_does_not_receive_the_coding_name():
    s,_=session_with(imported(gene='TET2-AS1,TET2',transcript='NR_126420.1, NM_001127208.3',
        coding='c.2815C>T',protein='p.Gln939Ter'))
    assert [(v.gene,v.coding,v.protein) for v in s.variants]==[
        ('TET2-AS1','',''),('TET2','c.2815C>T','p.Gln939Ter')]
    assert s.variants[1].raw['_annotation_group']['note']
    s.variants[0].raw['_values'].append('modified child')
    assert 'modified child' not in s.variants[1].raw['_values']


@pytest.mark.parametrize('changes',[
    dict(transcript='NM_1.1, NM_2.2',coding='c.1A>T',protein='p.Lys1Asn'),
    dict(transcript='NM_1.1',coding='c.1A>T',protein='p.Lys1Asn'),
    dict(alt='A,T'),
    dict(transcript='NM_1.1, NR_2.2',coding='c.1A>T, c.2A>T'),
    dict(transcript='NM_1.1, NR_2.2',coding='NM_1.1:c.1A>T, NR_2.2:c.2A>T'),
    dict(coding='NM_000314.8:c.19G>A, NM_001126049.2:c.-2001C>T'),
])
def test_ambiguous_mapping_cannot_bypass_review_through_genomic_normalization(changes):
    from solide.nomenclature import _normalization_query
    s,_=session_with(imported(**changes))
    assert 1<=len(s.variants)<=2
    for v in s.variants:
        assert not annotation_resolution_current(v)
        with pytest.raises(ValueError):_normalization_query(v)
        with pytest.raises(ValueError):query_record(v,'ClinVar')


def test_full_noncoding_name_remains_one_complete_reference():
    from solide.nomenclature import _normalization_query
    s,_=session_with(imported(transcript='NM_1.1, NR_2.2',
        coding='NM_1.1:c.1A>T, NR_2.2:n.2G>A',protein='p.Lys1Asn, '))
    assert len(s.variants)==2
    confirm(s.variants[1])
    assert query_record(s.variants[1],'ClinVar').hgvsc=='NR_2.2:n.2G>A'
    assert _normalization_query(s.variants[1])==('NR_2.2:n.2G>A','full_hgvs')


def test_ordinary_identical_rows_are_not_deduplicated_with_annotation_siblings(tmp_path):
    from solide.annotations import source_row_groups
    from solide.reporting import export_patient
    overlap=imported(coverage=300)
    ordinary=imported(gene='EGFR',transcript='NM_005228.5',coding='c.2573T>G',protein='p.Leu858Arg',source_row=21)
    duplicate=deepcopy(ordinary);duplicate.id='another-source-row';duplicate.source_row=22
    s=Session();a=add_analysis(s,ImportResult([overlap,ordinary,duplicate]),'00OUM00001','GRCh37','Ion Torrent','Childhood')
    assert len(s.variants)==4 and len(source_row_groups(s.variants))==3
    for v in s.variants:
        v.selected=True
        if v.raw.get('_annotation_group'):confirm(v)
    with closing(load_workbook(export_analysis(s,a.id,tmp_path).report)) as book:
        assert book['Rådata'].max_row==4
        assert [book['Rådata'].cell(row,1).value for row in range(2,5)]==['KLLN,PTEN','EGFR','EGFR']
    with closing(load_workbook(export_patient(s,a.patient,tmp_path))) as book:
        assert book['Raw data'].max_row==4
        assert book['Overview']['D4'].value==3
        assert book['Overview']['H4'].value==1
        links=[book['Quality'].cell(row,7).hyperlink.target for row in range(2,4)]
        assert links[0]==links[1] and 'A2' in links[0]


def test_gui_separates_purple_alternatives_and_searches_only_checked_child(tmp_path):
    import os
    os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
    from PyQt6.QtWidgets import QApplication
    from PyQt6.QtCore import Qt
    from solide.gui import MainWindow
    from solide.importing import load_file
    app=QApplication.instance() or QApplication([]);view=MainWindow()
    path=tmp_path/'00OUM00001.tsv'
    path.write_text('Type\tGenes\tTranscript\tCoding\tAmino Acid Change\tLocus\tRef\tAlt\tCoverage\n'
        'SNV\tKLLN,PTEN\tNM_001126049.2, NM_000314.8\tc.-2001C>T, c.19G>A\tp.?, p.Glu7Lys\tchr10:89624245\tG\tA\t600\n')
    try:
        result=load_file(path)
        assert len(result.variants)==1
        view.accept_import([(str(path),result,'')])
        assert view.variant_model.rowCount()==2
        assert view.import_model.item(0,3).text()=='1'
        assert '-1 duplicates' not in view.import_log.toPlainText()
        for row in range(2):
            assert view.variant_model.data(view.variant_model.index(row,2),Qt.ItemDataRole.BackgroundRole).name()=='#eadcf5'
            assert 'samme kilderad' in view.variant_model.data(view.variant_model.index(row,2),Qt.ItemDataRole.ToolTipRole)
        index=view.variant_model.index(1,0)
        assert view.variant_model.setData(index,Qt.CheckState.Checked,Qt.ItemDataRole.CheckStateRole)
        assert [v.gene for v in view.session.variants if v.selected]==['PTEN']
        jobs=[];view.launch=jobs.append
        for source,check in view.source_checks.items():check.setChecked(source=='ClinVar')
        view.start_queue()
        assert len(jobs)==1
        assert jobs[0].plan=={(view.session.variants[1].id,'ClinVar')}
        view.run_timer.stop()
    finally:view.close()
