from datetime import datetime
import zipfile
from openpyxl import load_workbook
from solide.models import Session,Variant
from solide.reporting import export_patient


def test_report_review_fields_and_original_raw_order(tmp_path):
    v=Variant(patient='DEMO',gene='EGFR',transcript='NM_005228.5',coding='c.2573T>G',protein='p.L858R',
        selected=True,af_percent=21.1,coverage=400,classification='Reviewed laboratory classification',
        report_decision='Include',reviewer='DEMO reviewer',raw={'Transcript':'NM_005228.5','Genes':'EGFR','Coding':'c.2573T>G',
        'Amino Acid Change':'p.L858R','Variant ID':'COSM000','ClinVar':'source annotation','Allele Frequency %':21.1})
    w=load_workbook(export_patient(Session(variants=[v]),'DEMO',tmp_path))
    overview=w['Overview'];headers={cell.value:cell.column for row in w['Evidence'] for cell in row if cell.value in {'Classification','Report decision','Reviewer'}}
    assert set(headers)=={'Classification','Report decision','Reviewer'}
    assert any(c.value=='Include' for row in overview for c in row)
    assert any(c.value=='Reviewed laboratory classification' for row in overview for c in row)
    assert list(w['Raw data'].values)[0][:7]==tuple(v.raw)
    assert 'SolideFindings' in overview.tables
    assert overview['A1'].value=='SOLIDE'


def test_exports_create_separate_versions_and_typed_dates(tmp_path):
    s=Session(variants=[Variant(patient='DEMO',gene='EGFR',selected=True)])
    a=export_patient(s,'DEMO',tmp_path);b=export_patient(s,'DEMO',tmp_path)
    assert a!=b and a.exists() and b.exists()
    w=load_workbook(a)
    assert any(isinstance(c.value,datetime) for row in w['Overview'] for c in row)


def test_stale_capture_and_full_mtbp_report_not_embedded(tmp_path):
    from PIL import Image
    image=tmp_path/'old.png';Image.new('RGB',(400,200),'green').save(image)
    v=Variant(patient='DEMO',gene='EGFR',selected=True)
    v.evidence['MTBP']={'database':'MTBP','status':'found','fingerprint':'old',
        'raw':{'screenshots':[{'path':str(image)}],'patient_report_screenshot':str(image)}}
    p=export_patient(Session(variants=[v]),'DEMO',tmp_path)
    with zipfile.ZipFile(p) as archive:assert not any(name.startswith('xl/media/') for name in archive.namelist())
    w=load_workbook(p)
    assert any('Outdated' in str(c.value) for sheet in w for row in sheet for c in row)


def test_tall_captures_are_segmented_without_overlapping_next_source(tmp_path):
    from PIL import Image,ImageDraw
    from solide.report_layout import put_image
    from openpyxl import Workbook
    image=Image.new('RGB',(1600,5000),'white');ImageDraw.Draw(image).rectangle((0,0,800,5000),fill='green')
    path=tmp_path/'tall.png';image.save(path)
    w=Workbook();sheet=w.active;end=put_image(sheet,path,2,'Capture')
    assert len(sheet._images)==4
    assert all(im.width==1050 and im.height<=900 for im in sheet._images)
    last=sheet._images[-1];row=int(last.anchor[1:])
    assert end>row+last.height/24
    assert image.size==(1600,5000)


def test_raw_text_formula_characters_preserved_as_strings(tmp_path):
    v=Variant(patient='DEMO',gene='EGFR',selected=True,raw={'A':'=1+1','B':'-literal','C':'@text'})
    w=load_workbook(export_patient(Session(variants=[v]),'DEMO',tmp_path))
    cells=[w['Raw data'].cell(2,column) for column in (1,2,3)]
    assert [cell.value for cell in cells]==['=1+1','-literal','@text']
    assert all(cell.data_type=='s' for cell in cells)


def test_overview_is_compact_and_links_unsearched_variants(tmp_path):
    v=Variant(patient='SOURCE',gene='EGFR',coding='c.2573T>G',selected=True)
    w=load_workbook(export_patient(Session(variants=[v]),'SOURCE',tmp_path))
    sheet=w['Overview']
    assert sheet.max_column==8
    assert all(isinstance(c.value,str) for c in sheet[5])
    assert sheet.row_dimensions[6].height<=72
    assert sheet['H6'].value=='Not run'
    assert sheet['B6'].hyperlink is not None
    assert sheet['B6'].hyperlink.target.startswith("#'Evidence'!")
    assert sheet['B6'].font.underline=='single'
    assert [s.title for s in w if s.sheet_state=='visible']==['Overview','Quality','Evidence','Raw data']
    assert w['Searches'].sheet_state=='hidden'


def test_evidence_navigation_preserves_long_notes_and_qc_statuses(tmp_path):
    note='A full reviewed assessment. '*180
    a=Variant(patient='SOURCE',gene='EGFR',coding='c.1A>T',selected=True,comment=note,reviewer='Reviewer',report_decision='Include')
    b=Variant(patient='SOURCE',gene='MET',kind='CNV',copy_number=1)
    c=Variant(patient='SOURCE',gene='ALK',kind='RNAExonVariant',call='ABSENT')
    w=load_workbook(export_patient(Session(variants=[a,b,c]),'SOURCE',tmp_path))
    evidence=w['Evidence']
    assert ''.join(str(row[1].value or '') for row in evidence if row[0].value in {'Assessment','Assessment (continued)'})==note
    assert {row[3].value for row in list(w['Quality'])[1:]}=={'Review','Report'}
    for sheet in w:
        for row in sheet:
            for cell in row:
                if cell.hyperlink and cell.hyperlink.target.startswith('#'):
                    target,location=cell.hyperlink.target[1:].split('!')
                    assert target.strip("'") in w.sheetnames
                    assert w[target.strip("'")][location].row<=w[target.strip("'")].max_row


def test_filtered_export_scope_is_explicit_without_changing_raw_columns(tmp_path):
    v=Variant(patient='SOURCE',gene='EGFR',selected=True,raw={'Gene':'EGFR',
        '_export_scope':{'filtered_export':True,'total_variant_count':3385,'exported_variant_count':8}})
    w=load_workbook(export_patient(Session(variants=[v]),'SOURCE',tmp_path))
    assert '8 exported / 3,385 total' in w['Overview']['H2'].value
    assert '_export_scope' not in [c.value for c in w['Raw data'][1]]
    assert w.properties.identifier=='solide:variant-review-report'


def test_multiline_assessment_splits_before_excel_row_height_clips_text(tmp_path):
    note='\n'.join(f'Assessment line {i}' for i in range(50))
    v=Variant(patient='SOURCE',gene='EGFR',selected=True,comment=note)
    w=load_workbook(export_patient(Session(variants=[v]),'SOURCE',tmp_path))
    rows=[row for row in w['Evidence'] if row[0].value in {'Assessment','Assessment (continued)'}]
    assert len(rows)>=3
    assert ''.join(row[1].value for row in rows)==note
    for row in rows:
        lines=str(row[1].value).count('\n')+1
        assert w['Evidence'].row_dimensions[row[0].row].height>=16*lines+10


def test_api_source_release_and_normalization_are_readable_without_setting_app_classification(tmp_path):
    v=Variant(patient='SOURCE',gene='BRCA1',selected=True,classification='Laboratory assessment')
    v.evidence['BRCA Exchange']={'database':'BRCA Exchange','status':'needs_review','fingerprint':v.fingerprint('Other'),
        'raw':{'release':{'id':76,'name':70},'classification':{'expert':'Source assertion','expert_date':'2026-03-08'}}}
    v.evidence['Mutalyzer']={'database':'Mutalyzer','status':'needs_review','fingerprint':v.fingerprint('Other'),
        'raw':{'query_basis':'full_hgvs','original':'NM_007294.4:c.68_69delAG',
            'normalization':{'normalized_description':'NM_007294.4:c.68_69del'}}}
    w=load_workbook(export_patient(Session(variants=[v]),'SOURCE',tmp_path))
    fields={row[0].value:row[1].value for row in w['Evidence'] if row[0].value}
    assert fields['Dataset release']=='{"id": 76, "name": 70}'
    assert fields['Normalized HGVS']=='NM_007294.4:c.68_69del'
    assert fields['Source classification']=='Source assertion'
    assert w['Overview']['F6'].value=='Laboratory assessment'
