from contextlib import closing
from openpyxl import load_workbook
from solide.models import Analysis,Session,Variant
from solide.reporting import export_analysis


def test_app_text_is_norwegian_but_source_assertions_and_raw_headers_are_unchanged(tmp_path):
    a=Analysis(id='a',patient='26OUM12345',profile='OPA')
    v=Variant(analysis_id='a',patient=a.patient,gene='EGFR',coding='c.1A>T',selected=True,
              raw={'_columns':['Gene','Comment'],'_values':['EGFR','Original English source value']})
    for source,status,text in [('Franklin','found','Likely Benign'),('MTBP','found','Putative functionally relevant variants'),
                               ('SpliceAI','not_applicable','Outside the ±100 bp intronic rule.')]:
        v.evidence[source]={'database':source,'status':status,'clinical_significance':text,
                            'fingerprint':v.fingerprint('Other'),'summary':text}
    # No screenshot: preserve the source text in hidden evidence without pretending it is a completed result.
    with closing(load_workbook(export_analysis(Session(analyses=[a],variants=[v]),'a',tmp_path).report)) as book:
        assert book.sheetnames==['Oversikt','EGFR','Kilder','Rådata']
        assert book['Oversikt']['A4'].value=='Opprettet'
        assert [c.value for c in book['Oversikt'][7]][:5]==['Gen','HGVSc','HGVSp','Kildefunn','Kommentar']
        assert book['Oversikt']['G8'].value=='Ikke kjørt'
        assert book['Oversikt']['M8'].value=='Bilde mangler'
        assert book['Rådata']['A1'].value=='Gene'
        assert book['Rådata']['B1'].value=='Comment'
        assert book['Rådata']['B2'].value=='Original English source value'
        values=[c.value for row in book['Kilder'] for c in row]
        assert 'Likely Benign' in values
        assert 'Putative functionally relevant variants' in values
        assert not any('Screenshot not available' in str(c.value) for row in book['EGFR'] for c in row)


def test_current_source_classifications_are_not_translated(tmp_path):
    from test_vpm_analysis_reports import make_session,capture,result
    from solide.evidence import batch_fingerprint
    s,rows=make_session(Variant(gene='EGFR',selected=True))
    v=rows[0];image=capture(tmp_path)
    v.evidence['Franklin']=result(v,image,clinical_significance='Benign')
    v.evidence['MTBP']=result(v,image,source='MTBP',clinical_significance='Putative functionally relevant variants',
        analysis_id=v.analysis_id,mtbp_batch=batch_fingerprint(s,v.patient,v.analysis_id))
    v.evidence['MTBP']['raw']['patient_report_screenshot']=str(image)
    with closing(load_workbook(export_analysis(s,'analysis-a',tmp_path).report)) as book:
        assert book['Oversikt']['M8'].value=='Benign'
        assert book['Oversikt']['L8'].value=='Putative functionally relevant variants'
