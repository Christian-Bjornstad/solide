from openpyxl import load_workbook
from solide.models import Variant, Session
from solide.reporting import export_patient

def test_patient_separation_and_qc_on_unselected(tmp_path):
    a=Variant(patient='DEMO-A',gene='EGFR',coding='c.1A>T',af_percent=21.1,selected=True,comment='=HYPERLINK("bad")')
    a.evidence['ClinVar']={'status':'error','summary':'Timeout','fingerprint':a.fingerprint('Other')}
    s=Session(variants=[a,Variant(patient='DEMO-A',gene='MET',kind='CNV',copy_number=0.9),
                        Variant(patient='DEMO-B',gene='PRIVATE-GENE',selected=True)])
    p=export_patient(s,'DEMO-A',tmp_path)
    w=load_workbook(p)
    cells=[c for sh in w for row in sh for c in row if c.value is not None]
    assert not any('PRIVATE-GENE' in str(c.value) for c in cells)
    assert any('Timeout' in str(c.value) for c in cells)
    assert any('MET'==c.value for row in w['Quality'] for c in row)
    assert not any(c.data_type=='f' for c in cells)
    assert any(c.value==21.1 for row in w['Overview'] for c in row)

def test_stale_evidence_is_visible(tmp_path):
    v=Variant(patient='DEMO',gene='MET',selected=True)
    v.evidence['ClinVar']={'status':'found','summary':'OLD','fingerprint':'old'}
    w=load_workbook(export_patient(Session(variants=[v]),'DEMO',tmp_path))
    assert any('Outdated' in str(c.value) for row in w['Overview'] for c in row)

def test_filename_cannot_escape_directory(tmp_path):
    patient='../../DEMO'
    p=export_patient(Session(variants=[Variant(patient=patient,gene='MET',selected=True)]),patient,tmp_path)
    assert p.parent==tmp_path

def test_report_embeds_current_variant_image_and_full_mtbp_report(tmp_path):
    import zipfile
    from PIL import Image
    from solide.evidence import batch_fingerprint
    image=tmp_path/'synthetic.png'
    Image.new('RGB',(640,400),(90,130,70)).save(image)
    v=Variant(patient='DEMO',gene='EGFR',selected=True,assembly='GRCh37')
    session=Session(variants=[v])
    v.evidence['MTBP']={'database':'MTBP','status':'found','fingerprint':v.fingerprint('Other'),
        'mtbp_batch':batch_fingerprint(session,'DEMO'),
        'raw':{'screenshots':[{'path':str(image),'label':'Synthetic evidence'}],
               'patient_report_screenshot':str(image)}}
    report=export_patient(session,'DEMO',tmp_path)
    with zipfile.ZipFile(report) as archive:
        assert len([n for n in archive.namelist() if n.startswith('xl/media/')])==2
    w=load_workbook(report)
    assert 'MTBP attachment 1' in w.sheetnames
