"""The exported attachment contains captures, not the app's review worksheet."""
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest

from test_vpm_analysis_reports import make_session, result, capture, load_workbook
from solide.models import Variant
from solide.reporting import export_analysis


@pytest.mark.parametrize('profile',['OPA','Ion Torrent'])
def test_clean_attachment_has_blank_notes_no_quality_details_or_internal_links(tmp_path,profile):
    session,rows=make_session(Variant(gene='EGFR',selected=True,comment='App-only assessment',
                                    classification='App classification'),profile=profile)
    rows[0].evidence['Franklin']=result(rows[0],capture(tmp_path),clinical_significance='Benign')
    path=export_analysis(session,'analysis-a',tmp_path).report
    with load_workbook(path) as book:
        assert book.sheetnames==['Overview','EGFR','Sources','Raw data']
        assert book['Overview']['D3'].value is None
        assert book['Overview']['G8'].value=='Benign'
        assert book['Overview']['G8'].hyperlink.target=='https://example.org/record'
        assert book['Overview']['A8'].hyperlink is None
        visible='\n'.join(str(c.value) for s in book if s.sheet_state=='visible' for row in s for c in row if c.value)
        assert all(x not in visible for x in ['Review match','Verified match','Laboratory assessment',
                   'Source details','Match assessment','Decision: Pending','App-only assessment','App classification'])
        assert len(book['EGFR']._images)==1
        assert not any(s.tables for s in book)
        assert not any(c.hyperlink and c.hyperlink.target.startswith('#') for s in book for row in s for c in row)
    with ZipFile(path) as archive:
        assert not any(n.startswith('xl/tables/') for n in archive.namelist())


def test_actual_source_labels_and_cosmic_id_replace_match_status(tmp_path):
    session,rows=make_session(Variant(gene='EGFR',selected=True))
    v=rows[0]
    image=capture(tmp_path)
    for source,classification in [('MTBP','Putative functionally relevant variants'),
                                  ('Franklin','Likely Benign'),('ClinVar','Pathogenic'),('OncoKB','Oncogenic')]:
        v.evidence[source]=result(v,image,source=source,clinical_significance=classification)
    # Without the batch binding MTBP must not be presented as a valid result.
    from solide.evidence import batch_fingerprint
    v.evidence['MTBP'].update(analysis_id=v.analysis_id,mtbp_batch=batch_fingerprint(session,'SAMPLE','analysis-a'))
    v.evidence['MTBP']['raw'].update(functional_evidence='Evidence A (Curated)',patient_report_screenshot=str(image))
    v.evidence['COSMIC']=result(v,image,source='COSMIC',clinical_significance='',accession='COSM6224',
                                url='https://cancer.sanger.ac.uk/cosmic/mutation/overview?id=6224')
    with load_workbook(export_analysis(session,'analysis-a',tmp_path).report) as book:
        row=book['Overview']
        assert row['F8'].value=='Putative functionally relevant variants – Evidence A (Curated)'
        assert row['G8'].value=='Likely Benign'
        assert row['H8'].value=='Pathogenic'
        assert row['I8'].value=='Oncogenic'
        assert row['J8'].value=='COSM6224'
        assert row['J8'].hyperlink.target.endswith('id=6224')


def test_no_evidence_and_stale_assertions_do_not_become_positive_results(tmp_path):
    session,rows=make_session(Variant(gene='EGFR',selected=True))
    rows[0].evidence['Franklin']=result(rows[0],clinical_significance='Benign',fingerprint='old')
    with load_workbook(export_analysis(session,'analysis-a',tmp_path).report) as book:
        assert book['Overview']['G8'].value=='Outdated'
        assert book['Overview']['G8'].hyperlink is None
        assert book['Overview']['J8'].value is None


def test_saved_views_have_one_valid_selection_for_each_frozen_pane(tmp_path):
    from openpyxl import Workbook
    from solide.reporting import _save_report_bundle
    book=Workbook();sheet=book.active
    for pane in ['C2','C5','C9','A4']:sheet.freeze_panes=pane
    path=tmp_path/'views.xlsx';_save_report_bundle(book,None,path,None)
    with ZipFile(path) as archive:
        ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        view=ET.fromstring(archive.read('xl/worksheets/sheet1.xml')).find('s:sheetViews/s:sheetView',ns)
        selections=view.findall('s:selection',ns)
        assert len(selections)==1
        assert selections[0].attrib=={'pane':'bottomLeft','activeCell':'A4','sqref':'A4'}
