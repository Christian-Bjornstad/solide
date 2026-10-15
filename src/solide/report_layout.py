"""Formatting and image placement for the generated Excel report."""
from __future__ import annotations
from io import BytesIO
from math import ceil
from pathlib import Path
import json
import math
import re
from PIL import Image as PillowImage
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment,Font,PatternFill,Border,Side
from openpyxl.worksheet.table import Table,TableStyleInfo
from openpyxl.utils import get_column_letter

GREEN='425B3D'
INK='243128'
PALE='F2F5EF'
WARN='FFF2D5'
FAIL='FCE5E3'


def safe_text(value):
    if isinstance(value,str):
        value=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',value)
        if value.startswith(('=','+','-','@')):value="'"+value
        return value[:32767]
    return value


def add_row(sheet,values):
    sheet.append([safe_text(value) for value in values])


def original_row(sheet,values):
    """Retain source text without allowing Excel to interpret formulas."""
    row=sheet.max_row+1
    for index,value in enumerate(values,1):
        if isinstance(value,str):
            value=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',value)[:32767]
        cell=sheet.cell(row,index,value)
        if isinstance(value,str):
            cell.data_type='s'
            cell.quotePrefix=value.startswith(('=','+','-','@'))


def title(sheet,text,width=8):
    sheet.sheet_view.showGridLines=False
    sheet.sheet_view.zoomScale=90
    sheet.sheet_properties.tabColor=GREEN
    sheet.merge_cells(start_row=1,start_column=1,end_row=1,end_column=width)
    cell=sheet.cell(1,1,safe_text(text));cell.font=Font(name='Calibri',size=16,bold=True,color=GREEN)
    cell.alignment=Alignment(vertical='center')
    for column in range(1,width+1):
        sheet.cell(1,column).border=Border(bottom=Side(style='thin',color=GREEN))
    sheet.row_dimensions[1].height=32


def style_table(sheet,header_row=1,name=None):
    sheet.freeze_panes=f'C{header_row+1}'
    sheet.sheet_view.showGridLines=False
    sheet.sheet_properties.tabColor=GREEN
    last=get_column_letter(sheet.max_column)
    sheet.auto_filter.ref=f'A{header_row}:{last}{sheet.max_row}'
    for cell in sheet[header_row]:
        cell.fill=PatternFill('solid',fgColor=GREEN)
        cell.font=Font(name='Calibri',size=11,color='FFFFFF',bold=True)
        cell.alignment=Alignment(wrap_text=True,vertical='center')
    sheet.row_dimensions[header_row].height=34
    for row in sheet.iter_rows(min_row=header_row+1):
        sheet.row_dimensions[row[0].row].height=36
        for cell in row:
            cell.font=Font(name='Calibri',size=11,color='176B94' if cell.hyperlink else INK,
                           underline='single' if cell.hyperlink else None)
            cell.alignment=Alignment(vertical='top',wrap_text=True)
            if cell.row%2==0:cell.fill=PatternFill('solid',fgColor=PALE)
    for column in range(1,sheet.max_column+1):
        width=max((len(str(sheet.cell(row,column).value or '')) for row in range(header_row,min(sheet.max_row,header_row+25)+1)),default=16)
        sheet.column_dimensions[get_column_letter(column)].width=min(48,max(16,width+2))
    if name and sheet.max_row>header_row:
        table=Table(displayName=name,ref=f'A{header_row}:{last}{sheet.max_row}')
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium4',showRowStripes=True)
        sheet.add_table(table)
    sheet.print_title_rows=f'{header_row}:{header_row}'
    sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
    sheet.sheet_properties.pageSetUpPr.fitToPage=True


def status_cell(cell,status):
    if status in {'Failed','Error','Timeout','Identity mismatch','Missing capture','Sign-in required','Ambiguous match'}:
        cell.fill=PatternFill('solid',fgColor=FAIL)
    elif status in {'Unknown','Review','Review required','Review match','Outdated','Pending','Review query','Recheck required','Retry / review needed'}:
        cell.fill=PatternFill('solid',fgColor=WARN)


def internal_link(cell,sheet,row=1,label='Open evidence'):
    cell.value=safe_text(label)
    cell.hyperlink=f"#'{sheet.replace(chr(39),chr(39)*2)}'!A{row}"
    cell.font=Font(name='Calibri',size=11,color='176B94',underline='single')
    cell.alignment=Alignment(wrap_text=True,vertical='top')


def put_image(sheet,path,row,label='',crop_top=0):
    """Embed a copy in full-width segments; keep the original capture unchanged."""
    path=Path(path)
    try:
        with PillowImage.open(path) as original:
            original.load();width,height=original.size
            crop_top=max(0,min(int(crop_top or 0),height-1))
            max_width=1050;scale=min(1,max_width/width)
            segment_height=max(1,int(900/scale))
            for number,top in enumerate(range(crop_top,height,segment_height),1):
                bottom=min(height,top+segment_height)
                segment=original.crop((0,top,width,bottom))
                buffer=BytesIO();segment.save(buffer,format='PNG');buffer.seek(0)
                image=Image(buffer);image.width*=scale;image.height*=scale
                if label:
                    sheet.cell(row,1,safe_text(f'{label} — part {number}' if height-crop_top>segment_height else label))
                    sheet.cell(row,1).font=Font(bold=True,color=GREEN,size=11);row+=1
                occupied=ceil(image.height/24)
                for index in range(row,row+occupied+1):sheet.row_dimensions[index].height=18
                sheet.add_image(image,f'A{row}');row+=occupied+2
            return row
    except (OSError,ValueError,PillowImage.DecompressionBombError):
        sheet.cell(row,1,'Screenshot could not be read.');return row+2


def mtbp_content_top(path):
    """Return a validated display crop; the original capture stays untouched."""
    path=Path(path);geometry=path.with_suffix('.geometry.json')
    if not path.is_file() or not geometry.is_file():return 0
    try:
        data=json.loads(geometry.read_text(encoding='utf8'))
        top=float(data.get('content_top') or 0);height=float(data.get('height') or 0)
        if top<=0:
            candidates=[]
            for item in data.get('rows') or []:
                if not isinstance(item,dict):continue
                box=item.get('section') or item.get('header') or item.get('row')
                if isinstance(box,dict):
                    value=float(box.get('y') or 0)
                    if math.isfinite(value) and value>0:candidates.append(value)
            top=min(candidates,default=0)
        if not math.isfinite(top) or not math.isfinite(height) or not 0<top<height:return 0
        with PillowImage.open(path) as original:
            pixel_top=round(top*original.height/height)
            return pixel_top if 0<pixel_top<original.height-1 else 0
    except (OSError,ValueError,TypeError,json.JSONDecodeError,AttributeError):return 0


# Archer/VPM patient_excel.py, pinned at c12a1b8 (layout unchanged since
# b5f50c6). Keep these reference styles separate from the legacy report styles.
VPM_COLORS={'navy':'163B5C','blue':'2F75B5','pale_blue':'EAF3FA',
            'gray':'F3F6F8','muted':'5E6A73','border':'CBD7E1',
            'white':'FFFFFF','pale_orange':'FFF3E8'}


def vpm_base_sheet(sheet,columns=12):
    sheet.sheet_view.showGridLines=False
    for column in range(1,columns+1):sheet.column_dimensions[get_column_letter(column)].width=13
    sheet.page_setup.orientation='landscape';sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
    sheet.sheet_properties.pageSetUpPr.fitToPage=True
    sheet.page_margins.left=sheet.page_margins.right=.25
    sheet.page_margins.top=sheet.page_margins.bottom=.4


def vpm_border():
    return Border(**{side:Side(style='thin',color=VPM_COLORS['border'])
                     for side in ('left','right','top','bottom')})


def vpm_title(sheet,text,columns=12):
    sheet.merge_cells(start_row=1,start_column=1,end_row=2,end_column=columns)
    cell=sheet.cell(1,1,safe_text(text))
    cell.font=Font(size=22,bold=True,color=VPM_COLORS['white'])
    cell.fill=PatternFill('solid',fgColor=VPM_COLORS['navy'])
    cell.alignment=Alignment(vertical='center',horizontal='left')


def vpm_section(sheet,row,text,columns=12,source=None):
    sheet.merge_cells(start_row=row,start_column=1,end_row=row,end_column=columns)
    cell=sheet.cell(row,1,safe_text(text))
    fills={'Franklin':'D9EAF7','OncoKB':'EAF5ED','MTBP':'FFF3E8'}
    cell.font=Font(size=13,bold=True,color=VPM_COLORS['navy'] if source else VPM_COLORS['white'])
    cell.fill=PatternFill('solid',fgColor=fills.get(source,VPM_COLORS['gray']) if source else VPM_COLORS['blue'])
    cell.alignment=Alignment(vertical='center')
    return row+1


def vpm_header(cell,text):
    cell.value=safe_text(text);cell.fill=PatternFill('solid',fgColor=VPM_COLORS['navy'])
    cell.font=Font(bold=True,color=VPM_COLORS['white'])
    cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    cell.border=vpm_border()


def vpm_info_row(sheet,row,label,value,end_column=12):
    sheet.merge_cells(start_row=row,start_column=1,end_row=row,end_column=2)
    if end_column>3:sheet.merge_cells(start_row=row,start_column=3,end_row=row,end_column=end_column)
    sheet.cell(row,1,safe_text(label));cell=sheet.cell(row,3,value)
    if isinstance(value,str):cell.data_type='s';cell.quotePrefix=value.startswith(('=','+','-','@'))
    sheet.cell(row,1).font=Font(bold=True,color=VPM_COLORS['navy'])
    sheet.cell(row,1).fill=PatternFill('solid',fgColor=VPM_COLORS['pale_blue'])
    sheet.cell(row,3).fill=PatternFill('solid',fgColor=VPM_COLORS['gray'])
    for cell in sheet[row]:cell.border=vpm_border()
    sheet.cell(row,3).alignment=Alignment(wrap_text=True,vertical='top')


def vpm_put_image(sheet,path,row,crop_top=0):
    """Match VPM's 1120px width and 38px rows without shrinking long captures.

    The uncropped image is embedded as the original bytes. MTBP geometry may
    crop a lossless in-memory display copy; the capture on disk is never edited.
    """
    path=Path(path)
    with PillowImage.open(path) as source:
        source.load();width,height=source.size
        crop_top=max(0,min(int(crop_top or 0),height-1))
        if crop_top:
            buffer=BytesIO();source.crop((0,crop_top,width,height)).save(buffer,format='PNG');buffer.seek(0)
            image=Image(buffer);height-=crop_top
        else:image=Image(str(path))
    scale=min(1,1120/max(1,width))
    image.width=max(1,int(width*scale));image.height=max(1,int(height*scale))
    sheet.add_image(image,f'A{row}')
    occupied=max(1,ceil(image.height/38))
    for current in range(row,row+occupied):sheet.row_dimensions[current].height=28.5
    return row+occupied
