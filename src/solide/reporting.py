"""Versioned Excel reports produced from the app's reviewed session."""
from __future__ import annotations
from datetime import datetime
from dataclasses import dataclass
import hashlib
import json
import os
from pathlib import Path
import re
import tempfile
import uuid
from openpyxl import Workbook
from openpyxl.styles import Font,Alignment,PatternFill
from .models import Session
from .quality import qc_flags,review_reasons
from .evidence import evidence_is_current,assess_evidence,valid_capture
from .report_layout import safe_text,add_row,original_row,style_table,title,status_cell,internal_link,put_image,mtbp_content_top,GREEN,INK,PALE
from .report_layout import (VPM_COLORS,vpm_base_sheet,vpm_border,vpm_title,vpm_section,
                            vpm_header,vpm_info_row,vpm_put_image)
from . import __version__
from .report_language import report_text,capture_caption
from .localtime import oslo_now, oslo_datetime, EXCEL_DATETIME_FORMAT
from .nomenclature import display_identity
from .annotations import source_row_groups,annotation_group,require_resolved_annotations

STATUS={'found':'Found','verified':'Verified','not_found':'Not found','not_applicable':'Not applicable',
    'needs_review':'Review required','error':'Error','timeout':'Timeout','identity_mismatch':'Identity mismatch',
    'partial_capture':'Partial capture','manual_required':'Manual search','login_required':'Sign-in required',
    'verification_required':'Review required','authentication_failed':'Sign-in failed','network_error':'Network error',
    'ambiguous_result':'Ambiguous match','invalid_query':'Review query','unsupported_query':'Review query',
    'submission_unknown':'Submission uncertain','rate_limited':'Rate limited','quota_exhausted':'Quota reached',
    'deferred':'Deferred','unauthorized':'Access required','token_required':'Access required',
    'layout_changed':'Layout changed','ambiguous':'Ambiguous match'}
SOURCE_ORDER=('Mutalyzer','SpliceAI','BRCA Exchange','OncoKB','COSMIC','ClinVar','MTBP','Franklin')


def source_classification(evidence):
    raw=evidence.get('raw',{})
    value=evidence.get('clinical_significance') or raw.get('expert_classification') or raw.get('classification') or ''
    if isinstance(value,dict):return value.get('expert','')
    return str(value)


def captured_time(value):
    try:return oslo_datetime(value).replace(tzinfo=None)
    except (ValueError,TypeError,OverflowError):return None


def ordered_evidence(variant):
    order={name:index for index,name in enumerate(SOURCE_ORDER)}
    return sorted(variant.evidence.items(),key=lambda pair:order.get(pair[0],99))


def original_value(v,*keys):
    return next((v.raw[key] for key in keys if v.raw.get(key) not in (None,'')),'')


def raw_values(v):
    columns=v.raw.get('_columns',[]);values=v.raw.get('_values',[])
    if not columns:return {key:value for key,value in v.raw.items() if not key.startswith('_')}
    result={};reserved={str(name).casefold() for name in columns if name}
    for index,name in enumerate(columns):
        label=str(name) if name else f'Unnamed column {index+1}'
        base=label;number=1
        while label.casefold() in {key.casefold() for key in result} or (not name and label.casefold() in reserved):
            number+=1;label=f'{base} ({number})'
        result[label]=values[index] if index<len(values) else None
    return result


def evidence_summary(v,session):
    assessments=[assess_evidence(v,e,session) for _,e in ordered_evidence(v)]
    if not assessments:return 'Not run'
    labels=[a.label for a in assessments]
    verified=labels.count('Verified match');matches=labels.count('Review match')
    if 'Identity mismatch' in labels:state='Identity mismatch'
    elif 'Outdated' in labels:state='Outdated'
    elif any(a.retryable for a in assessments):state='Retry / review needed'
    elif any(label=='Review required' for label in labels):state='Review required'
    elif verified:state='Verified match'
    elif matches:state='Review match'
    elif 'No match' in labels:state='No match'
    else:state='Not applicable'
    counts=[]
    if verified:counts.append(f'{verified} verified')
    if matches:counts.append(f'{matches} review match')
    return state+('\n'+', '.join(counts) if counts else '')


def variant_label(v):
    values=[v.corrected_hgvs or v.coding or v.protein]
    if v.protein and v.protein not in values:values.append(v.protein)
    if v.transcript:values.append(v.transcript)
    return '\n'.join(filter(None,values)) or 'Open variant'


def scope_note(variants):
    scopes={json.dumps(v.raw.get('_export_scope',{}),sort_keys=True) for v in variants
            if v.raw.get('_export_scope',{}).get('filtered_export')}
    notes=[]
    for scope in sorted(scopes):
        data=json.loads(scope);total=data.get('total_variant_count')
        exported=data.get('exported_variant_count')
        count=f'{exported:,} exported / {total:,} total' if total is not None else f'{exported:,} exported'
        notes.append(f'Filtered export: {count}. QC covers supplied rows only.')
    return '\n'.join(notes) or 'QC covers supplied rows only.'


def section_heading(sheet,row,label):
    sheet.merge_cells(start_row=row,start_column=1,end_row=row,end_column=8)
    cell=sheet.cell(row,1,safe_text(label));cell.font=Font(name='Calibri',size=13,bold=True,color='FFFFFF')
    for column in range(1,9):sheet.cell(row,column).fill=PatternFill('solid',fgColor=GREEN)
    cell.alignment=Alignment(vertical='center',wrap_text=True);sheet.row_dimensions[row].height=32
    return row+1


def readable_chunks(text):
    start=0;lines=1;column=0
    for index,char in enumerate(text):
        if char=='\n':lines+=1;column=0
        else:
            column+=1
            if column>100:lines+=1;column=1
        if lines>=20 or index-start>=899:
            yield text[start:index+1];start=index+1;lines=1;column=0
    if start<len(text):yield text[start:]


def service_fields(source,data):
    basis={'full_hgvs':'Full versioned HGVS','genomic_grch37':'GRCh37 genomic alleles',
           'gene_cdna':'Gene / cDNA','gene_protein_or_identifier':'Gene / protein or identifier'}.get(data.get('query_basis'))
    fields=[('Query basis',basis)]
    if source=='Mutalyzer':
        candidate=data.get('genomic_normalization',{})
        fields.extend([('Requested HGVS',data.get('original')),('Normalized HGVS',data.get('normalization',{}).get('normalized_description')),
            ('Target transcript',data.get('mane_target')),('Mapped HGVS',data.get('mapping',{}).get('mapped_description')),
            ('Mapping limitation',data.get('mapping_issue')),('Genomic query',candidate.get('original')),
            ('Genomic candidate',candidate.get('normalization',{}).get('normalized_description')),
            ('Candidate status','Review required' if candidate else data.get('genomic_normalization_issue'))])
    elif source=='SpliceAI':
        fields.extend([('Requested allele',data.get('query')),('Scored allele',data.get('scored_variant')),
            ('Prediction settings',f"GRCh{data['hg']}; distance={data['distance']}; mask={data['mask']}"
                if all(key in data for key in ('hg','distance','mask')) else None)])
    elif source=='BRCA Exchange':
        release=data.get('release') or {}
        classification=data.get('classification');classification=classification if isinstance(classification,dict) else {}
        fields.extend([('Dataset release',json.dumps(release,ensure_ascii=False,default=str)),
            ('Expert evaluated',classification.get('expert_date'))])
    return [(label,value) for label,value in fields if value not in (None,'')]


def evidence_field(sheet,row,label,value):
    """Split long notes into readable rows without discarding their text."""
    text=str(value or '')
    chunks=list(readable_chunks(text)) or ['']
    for index,chunk in enumerate(chunks):
        sheet.cell(row,1,label if index==0 else label+' (continued)').font=Font(name='Calibri',size=11,bold=True,color=INK)
        sheet.cell(row,1).alignment=Alignment(wrap_text=True,vertical='top')
        sheet.merge_cells(start_row=row,start_column=2,end_row=row,end_column=8)
        content=value if isinstance(value,datetime) else safe_text(chunk)
        cell=sheet.cell(row,2,content);cell.font=Font(name='Calibri',size=11,color=INK)
        cell.alignment=Alignment(wrap_text=True,vertical='top')
        lines=max(sum(max(1,(len(line)+99)//100) for line in chunk.split('\n')),
                  (len(str(sheet.cell(row,1).value))+17)//18)
        sheet.row_dimensions[row].height=max(30,min(390,16*lines+10))
        if isinstance(value,datetime):cell.number_format=EXCEL_DATETIME_FORMAT
        if label=='Source URL' and text.startswith(('https://','http://')):
            cell.hyperlink=text;cell.style='Hyperlink';cell.alignment=Alignment(wrap_text=True,vertical='top')
        row+=1
    return row


def export_patient(session:Session,patient:str,directory:Path)->Path:
    analyses=[a for a in getattr(session,'analyses',[]) if a.patient==patient]
    patient_rows=[v for v in session.variants if v.patient==patient]
    bound_ids={v.analysis_id for v in patient_rows if getattr(v,'analysis_id','')}
    if len(analyses)>1 or len(bound_ids)>1 or (analyses and any(v.analysis_id!=analyses[0].id for v in patient_rows)):
        raise ValueError('This sample has multiple analyses. Choose one analysis before exporting.')
    variants=[v for v in session.variants if v.patient==patient]
    if not variants:raise ValueError('No rows for this patient.')
    original_groups=source_row_groups(variants)
    selected=[v for v in variants if v.selected]
    for variant in selected:require_resolved_annotations(variant)
    directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    stem=re.sub(r'[^\w.-]','_',patient).strip('.')[:70] or 'patient'
    generated=oslo_now();report_id=uuid.uuid4().hex[:8]
    path=directory/f'{stem}_{hashlib.sha256(patient.encode()).hexdigest()[:6]}_Solide_{generated:%Y%m%d_%H%M%S}_{report_id}.xlsx'
    w=Workbook();w.properties.identifier='solide:variant-review-report'
    overview=w.active;overview.title='Overview';title(overview,'SOLIDE',8)
    add_row(overview,['Patient / sample',patient,'MTBP tissue',session.tissue(patient),'Report ID',report_id,'Export scope',scope_note(variants)])
    add_row(overview,['Generated (Oslo)',generated.replace(tzinfo=None),'App version',__version__,'Report format',3])
    assemblies=', '.join(sorted({v.assembly for v in variants}))
    add_row(overview,['Assembly',assemblies,'Imported rows',len(original_groups),'Selected variants',len(selected),'Failed QC rows',
        sum(any(f.status=='Failed' for f in qc_flags(group[0])) for group in original_groups)])
    headers=['Gene','Variant','AF (%)','Coverage','Quality','Classification','Report decision','Evidence']
    add_row(overview,headers)
    for v in selected:
        flags=qc_flags(v)
        quality='; '.join(dict.fromkeys(flag.status for flag in flags)) or 'No row flags'
        add_row(overview,[v.gene,variant_label(v),v.af_percent,v.coverage,quality,
            v.classification,v.report_decision,evidence_summary(v,session)])
    style_table(overview,5,'SolideFindings')
    widths={'A':12,'B':35,'C':9,'D':11,'E':14,'F':22,'G':17,'H':24}
    for column,width in widths.items():overview.column_dimensions[column].width=width
    overview.freeze_panes='C6';overview.sheet_view.zoomScale=90
    overview.cell(3,2).number_format=EXCEL_DATETIME_FORMAT
    for row,v in enumerate(selected,6):
        overview.cell(row,3).number_format='0.00" %"';overview.cell(row,4).number_format='#,##0'
        lines=max(sum(max(1,(len(line)+int(widths[column])-1)//int(widths[column]))
                      for line in str(overview[f'{column}{row}'].value or '').split('\n')) for column in ('B','F','H'))
        overview.row_dimensions[row].height=max(54,min(120,16*lines+10))
        flags=qc_flags(v)
        status_cell(overview.cell(row,7),v.report_decision)
        status_cell(overview.cell(row,5),'Failed' if any(f.status=='Failed' for f in flags) else overview.cell(row,5).value)
        status_cell(overview.cell(row,8),str(overview.cell(row,8).value).split('\n')[0])
    for row in (2,3,4):
        overview.row_dimensions[row].height=60 if row==2 else 34
        for cell in overview[row]:cell.font=Font(name='Calibri',size=11,color=INK);cell.alignment=Alignment(wrap_text=True,vertical='center')
    overview.print_area=f'A1:H{overview.max_row}'

    qc=w.create_sheet('Quality')
    raw_rows={v.id:index for index,group in enumerate(original_groups,2) for v in group}
    add_row(qc,['Gene','Variant','Category','Status','Value','Reason','Source row','Source file','Row ID'])
    for v in variants:
        for flag in qc_flags(v):
            value=v.coverage if flag.category=='Coverage' else v.copy_number if flag.category.startswith('CNV') else v.call
            add_row(qc,[v.gene,v.corrected_hgvs or v.coding or v.protein,flag.category,flag.status,value,flag.message,v.source_row,Path(v.source_file).name,v.id])
            internal_link(qc.cell(qc.max_row,7),'Raw data',raw_rows[v.id],label=str(v.source_row))
    style_table(qc,name='SolideQuality')
    for letter,width in {'A':12,'B':32,'C':24,'D':14,'E':14,'F':48,'G':12}.items():qc.column_dimensions[letter].width=width
    for letter in ('H','I'):qc.column_dimensions[letter].hidden=True
    for row in range(2,qc.max_row+1):
        status_cell(qc.cell(row,4),qc.cell(row,4).value);qc.row_dimensions[row].height=54
    qc.print_area=f'A1:G{qc.max_row}'

    detail=w.create_sheet('Evidence');title(detail,'EVIDENCE',8)
    add_row(detail,['Patient / sample',patient,'Assembly',assemblies])
    detail.cell(3,1,'App assessments and source assertions are recorded separately.');detail.merge_cells('A3:H3')
    internal_link(detail.cell(4,1),'Overview',label='Back to findings')
    add_row(detail,['Gene','Variant','Classification','Report decision','Reviewer','Reviewed at (Oslo)','Quality','Evidence'])
    for v in selected:add_row(detail,[v.gene,v.corrected_hgvs or v.coding or v.protein,v.classification,v.report_decision,
        v.reviewer,captured_time(v.reviewed_at),'; '.join(f.status for f in qc_flags(v)) or 'No row flags','Open evidence'])
    style_table(detail,5,'SolideEvidenceIndex')
    for column in 'ABCDEFGH':detail.column_dimensions[column].width=18
    detail.freeze_panes='A6';detail.sheet_view.zoomScale=90
    for row in range(6,detail.max_row+1):
        detail.row_dimensions[row].height=48;detail.cell(row,6).number_format=EXCEL_DATETIME_FORMAT
    cursor=detail.max_row+3;anchors={};source_anchors={};full_reports={};full_links=[]
    for index,v in enumerate(selected,6):
        anchors[v.id]=cursor
        internal_link(overview.cell(index,2),'Evidence',cursor,label=variant_label(v))
        internal_link(overview.cell(index,8),'Evidence',cursor,label=evidence_summary(v,session))
        internal_link(detail.cell(index,2),'Evidence',cursor,label=v.corrected_hgvs or v.coding or v.protein or 'Open variant')
        internal_link(detail.cell(index,8),'Evidence',cursor)
        cursor=section_heading(detail,cursor,f'{v.gene}  {v.corrected_hgvs or v.coding or v.protein}')
        internal_link(detail.cell(cursor,1),'Overview',index,label='Back to finding')
        internal_link(detail.cell(cursor,3),'Evidence',5,label='Evidence index');cursor+=1
        fields=[('Original HGVS',':'.join(filter(None,[v.transcript,v.coding]))),('Protein',v.protein),
            ('Reviewed HGVS',v.corrected_hgvs),('Genomic identity',' | '.join(filter(None,[v.assembly,v.locus,v.ref,v.alt]))),
            ('Reviewed allele',v.controlled_genomic),('Identity review','Approved' if v.nomenclature_verified else 'Not approved'),
            ('Type / call',' / '.join(filter(None,[v.kind,v.call]))),('AF / coverage',f'{v.af_percent if v.af_percent is not None else "Unknown"}% / {v.coverage if v.coverage is not None else "Unknown"}'),
            ('Imported variant ID',v.variant_id or original_value(v,'Variant ID')),('Imported ClinVar',original_value(v,'ClinVar')),
            ('Classification',v.classification),('Report decision',v.report_decision),('Reviewer',v.reviewer),
            ('Reviewed at (Oslo)',captured_time(v.reviewed_at)),('Assessment',v.comment),
            ('Required checks',', '.join(review_reasons(v))),('Quality', '\n'.join(f'{f.category}: {f.status} — {f.message}' for f in qc_flags(v)) or 'No row flags'),
            ('Source provenance',f'{Path(v.source_file).name} | sheet {v.raw.get("_sheet", "")} | row {v.source_row}'),
            ('Source SHA256',v.source_hash),('Row ID',v.id)]
        for label,value in fields:
            if value not in (None,''):cursor=evidence_field(detail,cursor,label,value)
        internal_link(detail.cell(cursor,1),'Raw data',raw_rows[v.id],label='Open original row');cursor+=2
        if not v.evidence:cursor=evidence_field(detail,cursor,'Search status','Not run')
        for source,e in ordered_evidence(v):
            source_anchors[v.id,source]=cursor
            assessment=assess_evidence(v,e,session);current=evidence_is_current(v,e,session.tissue(patient),session)
            cursor=section_heading(detail,cursor,source)
            cursor=evidence_field(detail,cursor,'Match assessment',assessment.label)
            status_cell(detail.cell(cursor-1,2),assessment.label)
            for label,value in [('Source result',STATUS.get(e.get('status'),e.get('status','Unknown'))),
                ('Source classification',source_classification(e)),('Accession',e.get('accession','')),
                ('Captured at (Oslo)',captured_time(e.get('captured_at',''))),('Source URL',e.get('url','')),('Summary',e.get('summary',''))]:
                if value not in (None,''):cursor=evidence_field(detail,cursor,label,value)
            data=e.get('raw',{});images=data.get('screenshots',[]) or ([{'path':data['screenshot'],'label':source}] if data.get('screenshot') else [])
            for label,value in service_fields(source,data):cursor=evidence_field(detail,cursor,label,value)
            if not current:
                cursor=evidence_field(detail,cursor,'Freshness','Outdated evidence: rerun before reporting.');cursor+=2;continue
            for info in images:
                image_path=info.get('path','')
                if valid_capture(image_path):cursor=put_image(detail,image_path,cursor,info.get('label',source))
                else:cursor=evidence_field(detail,cursor,'Capture','Screenshot missing or unreadable: capture again.')
            full=data.get('patient_report_screenshot')
            if full:
                full_reports.setdefault(full,e);full_links.append((cursor,full))
                detail.cell(cursor,1,'Open full MTBP report');cursor+=2
            cursor+=2
        cursor+=2
    full_anchors={}
    for number,(full,e) in enumerate(full_reports.items(),1):
        full_anchors[full]=cursor;cursor=section_heading(detail,cursor,f'MTBP full report {number}')
        internal_link(detail.cell(cursor,1),'Evidence',5,label='Evidence index');cursor+=1
        cursor=evidence_field(detail,cursor,'Captured at (Oslo)',captured_time(e.get('captured_at','')))
        if valid_capture(full):cursor=put_image(detail,full,cursor,'MTBP full report')
        else:cursor=evidence_field(detail,cursor,'Capture','Full report screenshot missing or unreadable. Capture again.')
        cursor+=2
    for row,full in full_links:internal_link(detail.cell(row,1),'Evidence',full_anchors[full],label='Open full MTBP report')
    detail.print_area=f'A1:H{max(detail.max_row,cursor-1)}'

    searches=w.create_sheet('Searches')
    add_row(searches,['Gene','Variant','Source','Result','Match assessment','Classification from source','Accession',
        'Captured at (Oslo)','Tissue','Source URL','Summary','Evidence','Row ID','JSON response'])
    for v in selected:
        for source,e in ordered_evidence(v):
            assessment=assess_evidence(v,e,session)
            add_row(searches,[v.gene,v.corrected_hgvs or v.coding or v.protein,source,STATUS.get(e.get('status'),e.get('status','Unknown')),
                assessment.label,source_classification(e),e.get('accession',''),captured_time(e.get('captured_at','')),
                e.get('tissue',''),e.get('url',''),e.get('summary',''),'Open evidence',v.id,
                json.dumps(e.get('raw',{}),ensure_ascii=False,default=str)])
            row=searches.max_row;searches.cell(row,8).number_format=EXCEL_DATETIME_FORMAT
            url=e.get('url','')
            if isinstance(url,str) and url.startswith(('https://','http://')):searches.cell(row,10).hyperlink=url
            internal_link(searches.cell(row,12),'Evidence',source_anchors[v.id,source])
    style_table(searches,name='SolideSearches')
    for letter,width in {'B':28,'E':24,'F':40,'J':55,'K':70,'L':20}.items():searches.column_dimensions[letter].width=width
    searches.column_dimensions['N'].hidden=True;searches.sheet_state='hidden'
    for row in range(2,searches.max_row+1):
        status_cell(searches.cell(row,5),searches.cell(row,5).value)
        searches.row_dimensions[row].height=max(48,min(180,18*(str(searches.cell(row,11).value or '').count('\n')+2)))

    raw=w.create_sheet('Raw data')
    # Preserve A–L and every subsequent exported column in source order.
    originals=[raw_values(group[0]) for group in original_groups]
    keys=list(dict.fromkeys(key for original in originals for key in original))
    technical=['Solide Row ID','Selected','Source row','Source file','Source sheet','Source SHA256','Assembly']
    occupied={key.casefold() for key in keys}
    for index,name in enumerate(technical):
        while name.casefold() in occupied:name+=' (provenance)'
        technical[index]=name;occupied.add(name.casefold())
    add_row(raw,[*keys,*technical])
    for group,original in zip(original_groups,originals):
        v=group[0]
        original_row(raw,[*[original.get(k) for k in keys],'; '.join(child.id for child in group),any(child.selected for child in group),v.source_row,Path(v.source_file).name,v.raw.get('_sheet',''),v.source_hash,v.assembly])
    style_table(raw,name='SolideRawData');raw.freeze_panes='C2'

    fd,tmp=tempfile.mkstemp(dir=directory,prefix='.solide-',suffix='.xlsx');os.close(fd)
    try:w.save(tmp);os.replace(tmp,path)
    finally:
        if os.path.exists(tmp):os.unlink(tmp)
        w.close()
    return path


@dataclass(frozen=True)
class ReportBundle:
    report: Path
    attachment: Path | None = None

    @property
    def paths(self) -> tuple[Path, ...]:
        return (self.report,self.attachment) if self.attachment else (self.report,)


_CAPTURE_BLOCKERS={'identity_mismatch','ambiguous','ambiguous_result','layout_changed',
    'login_required','authentication_failed','error','network_error','timeout','partial_capture'}


def _public_source_url(source,value):
    if not isinstance(value,str) or not value.startswith(('https://','http://')):return ''
    if source=='MTBP' and re.search(r'/(patients|analysisQueue)/',value,re.IGNORECASE):return ''
    return value


def _capture_export_allowed(variant,item,scoped):
    assessment=assess_evidence(variant,item,scoped)
    return (evidence_is_current(variant,item,scoped.tissue_for(variant),scoped)
            and item.get('status') not in _CAPTURE_BLOCKERS
            and item.get('raw',{}).get('nomenclature_assessment',{}).get('status')!='Conflict'
            and assessment.label not in {'Identity mismatch','Outdated','Ambiguous match','Missing capture',
                                         'Failed','Timeout','Layout changed','Partial capture','Sign-in required','Recheck required'})


def _literal_chunks(sheet,row,column,value):
    """Write a full string in adjacent literal cells, preserving exact joins."""
    from openpyxl.utils import get_column_letter
    text=str(value or '')
    if re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',text):
        raise ValueError('A source response contains characters Excel cannot store losslessly.')
    chunks=[text[index:index+32767] for index in range(0,len(text),32767)] or ['']
    for offset,chunk in enumerate(chunks):
        current=column+offset
        if current>16384:raise ValueError('A source response exceeds the available Excel storage columns.')
        cell=sheet.cell(row,current,chunk);cell.data_type='s';cell.quotePrefix=chunk.startswith(('=','+','-','@'))
        sheet.column_dimensions[get_column_letter(current)].hidden=True
    return len(chunks)


def _text_height(values,widths,minimum=36):
    lines=max((sum(max(1,(len(line)+width-1)//width) for line in str(value or '').split('\n'))
               for value,width in zip(values,widths)),default=1)
    return max(minimum,min(409,16*lines+10))


def _exported_response(source,value):
    """Exclude ephemeral private MTBP report URLs from exported response text."""
    if isinstance(value,dict):return {key:_exported_response(source,item) for key,item in value.items()}
    if isinstance(value,list):return [_exported_response(source,item) for item in value]
    if source=='MTBP' and isinstance(value,str):
        return re.sub(r'https?://mtbp\.org/(?:patients|analysisQueue)/[^\s\"<>]*',
                      '[local MTBP evidence]',value,flags=re.IGNORECASE)
    return value


def _relative_target(path,directory):
    try:return Path(os.path.relpath(Path(path).resolve(),Path(directory).resolve())).as_posix()
    except ValueError:return Path(path).resolve().as_uri()


def _link(cell,target,label):
    cell.value=safe_text(label);cell.hyperlink=target
    cell.font=Font(name='Calibri',size=11,color='176B94',underline='single')
    cell.alignment=Alignment(wrap_text=True,vertical='top')


def _workbook_marker(workbook):
    workbook.properties.identifier='solide:variant-review-report'
    workbook.properties.version='4'


def _base_compact(sheet,widths):
    sheet.sheet_view.showGridLines=False;sheet.sheet_view.zoomScale=90
    sheet.sheet_properties.tabColor=GREEN
    for column,width in enumerate(widths,1):
        from openpyxl.utils import get_column_letter
        sheet.column_dimensions[get_column_letter(column)].width=width
    sheet.page_setup.orientation='landscape';sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
    sheet.sheet_properties.pageSetUpPr.fitToPage=True


def _table_region(sheet,row,headers,records,widths,name):
    """Style only this table so later notes do not enter its filter/sort range."""
    from openpyxl.worksheet.table import Table,TableStyleInfo
    from openpyxl.utils import get_column_letter
    for column,value in enumerate(headers,1):
        cell=sheet.cell(row,column,safe_text(value));cell.font=Font(name='Calibri',size=11,bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor=GREEN);cell.alignment=Alignment(wrap_text=True,vertical='center')
    sheet.row_dimensions[row].height=32
    for offset,values in enumerate(records,1):
        current=row+offset;sheet.row_dimensions[current].height=36
        for column,value in enumerate(values,1):
            cell=sheet.cell(current,column,safe_text(value));cell.font=Font(name='Calibri',size=11,color=INK)
            cell.alignment=Alignment(wrap_text=True,vertical='top')
            if offset%2:cell.fill=PatternFill('solid',fgColor=PALE)
    _base_compact(sheet,widths)
    last=row+len(records)
    if records:
        table=Table(displayName=name,ref=f'A{row}:{get_column_letter(len(headers))}{last}')
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium4',showRowStripes=True);sheet.add_table(table)
    sheet.freeze_panes=f'C{row+1}';sheet.print_title_rows=f'{row}:{row}'
    return last+1


def _main_findings(workbook,analysis,variants,generated):
    selected=[variant for variant in variants if variant.selected]
    sheet=workbook.active;sheet.title='Findings';title(sheet,f'{analysis.profile} variant review',8)
    add_row(sheet,['Sample / dataset',analysis.patient,'Profile',analysis.profile,'Assay',analysis.assay])
    pending=sum(variant.report_decision=='Pending' for variant in selected)
    add_row(sheet,['Generated (Oslo)',generated,'Genome',analysis.assembly,'Review state',
                   f'Draft: {pending} pending' if pending else 'Recorded app assessments'])
    sheet.cell(3,2).number_format=EXCEL_DATETIME_FORMAT
    add_row(sheet,['Selected variants',len(selected)])
    records=[[variant.gene,variant.transcript,variant.corrected_hgvs or variant.coding,variant.protein,
              variant.af_percent,variant.coverage,variant.classification,variant.report_decision] for variant in selected]
    end=_table_region(sheet,5,['Gene','Source transcript','Reviewed / source variant','Protein','AF (%)','Coverage','Lab classification','Decision'],
        records,[13,21,29,25,9,11,24,13],'SolideAnalysisFindings')
    for row,variant in enumerate(selected,6):
        sheet.cell(row,5).number_format='0.00" %"';sheet.cell(row,6).number_format='#,##0'
        if variant.coverage is not None and variant.coverage<500:status_cell(sheet.cell(row,6),'Failed')
        # Most laboratory findings fit in two lines; unusually long identities
        # retain enough height instead of silently clipping their nomenclature.
        values=[variant.gene,variant.transcript,variant.corrected_hgvs or variant.coding,variant.protein,variant.classification]
        widths=[13,21,29,25,24]
        lines=max(sum(max(1,(len(line)+width-1)//width) for line in str(value or '').split('\n'))
                  for value,width in zip(values,widths))
        sheet.row_dimensions[row].height=max(36,min(390,16*lines+8))
    for row in (2,3,4):
        sheet.row_dimensions[row].height=_text_height([cell.value for cell in sheet[row]],
            [13,21,29,25,9,11,24,13],36) if row<4 else 25
        for cell in sheet[row]:cell.font=Font(name='Calibri',size=11,color=INK);cell.alignment=Alignment(wrap_text=True,vertical='center')
    notes=[variant for variant in selected if variant.comment or variant.reviewer or variant.reviewed_at
           or len(variant.classification)>48]
    if notes:
        end=section_heading(sheet,end+2,'Laboratory assessments')
        for variant in notes:
            end=evidence_field(sheet,end,'Variant',': '.join(filter(None,[variant.gene,variant.corrected_hgvs or variant.coding or variant.protein])))
            for label,value in [('Lab classification',variant.classification),('Decision',variant.report_decision),
                                ('Reviewer',variant.reviewer),('Reviewed at (Oslo)',captured_time(variant.reviewed_at)),
                                ('Assessment',variant.comment)]:
                if value not in (None,''):end=evidence_field(sheet,end,label,value)
            end+=1
    sheet.print_area=f'A1:H{max(sheet.max_row,end-1)}'


def _quality_sections(workbook,analysis,variants):
    sheet=workbook.create_sheet('Quality');title(sheet,f'{analysis.profile} quality review',7)
    note=scope_note(variants)
    sheet.merge_cells('A2:G2');sheet.cell(2,1,note).alignment=Alignment(wrap_text=True,vertical='top')
    sheet.row_dimensions[2].height=44
    row=4;number=0
    categories=[('Row coverage',lambda variant: variant.coverage is not None or variant.kind.replace(' ','').lower()
                 in {'snp','snv','indel','del','ins','mnp','complex'}, {'Coverage'}),
                ('CNV',lambda variant: variant.kind.replace(' ','').lower()=='cnv',{'CNV','CNV review'})]
    if analysis.profile=='OPA':
        categories.extend([('RNA expression imbalance',lambda variant: variant.kind.replace(' ','').lower()=='rnaexontiles',{'Expression imbalance'}),
                           ('RNA exon variants',lambda variant: variant.kind.replace(' ','').lower()=='rnaexonvariant',{'RNAExonVariant'})])
    widths=[13,30,16,12,43,12,18]
    for heading,predicate,names in categories:
        row=section_heading(sheet,row,heading)
        applicable=[variant for variant in variants if predicate(variant)]
        records=[]
        for variant in applicable:
            for flag in qc_flags(variant):
                if flag.category in names:
                    value=variant.coverage if flag.category=='Coverage' else variant.copy_number if flag.category.startswith('CNV') else variant.call
                    records.append([variant.gene,variant.corrected_hgvs or variant.coding or variant.protein,flag.status,value,
                                    flag.message,variant.source_row,variant.kind])
        if records:
            number+=1;first=row
            row=_table_region(sheet,row,['Gene','Variant','Status','Value','Reason','Source row','Type'],records,widths,
                              f'SolideAnalysisQuality{number}')
            for current,values in enumerate(records,first+1):
                status_cell(sheet.cell(current,3),values[2]);sheet.row_dimensions[current].height=_text_height(values,widths,44)
        else:
            row=evidence_field(sheet,row,'Status','No exceptions in supplied rows' if applicable else 'Not supplied')
        row+=2
    if analysis.profile!='OPA':
        row=section_heading(sheet,row,'Identity and nomenclature checks')
        checks=[[variant.gene,variant.corrected_hgvs or variant.coding or variant.protein,'Review','',
                 ', '.join(review_reasons(variant)),variant.source_row,variant.kind]
                for variant in variants if review_reasons(variant)]
        if checks:
            first=row
            row=_table_region(sheet,row,['Gene','Variant','Status','Value','Required checks','Source row','Type'],
                              checks,widths,'SolideAnalysisIdentityChecks')
            for current,values in enumerate(checks,first+1):
                sheet.row_dimensions[current].height=_text_height(values,widths,44)
        else:row=evidence_field(sheet,row,'Status','No additional identity checks identified')
        row=evidence_field(sheet,row+1,'Coverage scope','Variant row coverage does not establish whole-gene or run/chip coverage. Run/chip metrics: Not supplied.')
    # Unknown/Report/Review flags from unexpected record types remain visible,
    # even when their category does not belong to the chosen profile's sections.
    known=set().union(*(names for _,_,names in categories))
    extras=[[variant.gene,variant.corrected_hgvs or variant.coding or variant.protein,flag.status,
             flag.category,flag.message,variant.source_row,variant.kind]
            for variant in variants for flag in qc_flags(variant) if flag.category not in known]
    if extras:
        row=section_heading(sheet,row+1,'Additional supplied quality flags')
        row=_table_region(sheet,row,['Gene','Variant','Status','Category','Reason','Source row','Type'],extras,widths,
                          'SolideAnalysisAdditionalQuality')
    _base_compact(sheet,widths);sheet.freeze_panes='A4';sheet.print_area=f'A1:H{max(sheet.max_row,row)}'


def _raw_analysis_sheet(workbook,variants,analysis,*,norwegian=False):
    groups=source_row_groups(variants)
    sheet=workbook.create_sheet('Rådata' if norwegian else 'Raw data');originals=[raw_values(group[0]) for group in groups]
    keys=list(dict.fromkeys(key for values in originals for key in values))
    technical=['Solide Row ID','Selected','Source row','Source file','Source sheet','Assembly','Analysis ID','Source SHA256']
    if norwegian:technical=[report_text(label) for label in technical]
    occupied={key.casefold() for key in keys}
    for index,name in enumerate(technical):
        while name.casefold() in occupied:name+=' (kildedata)' if norwegian else ' (provenance)'
        technical[index]=name;occupied.add(name.casefold())
    add_row(sheet,[*keys,*technical])
    for group,values in zip(groups,originals):
        variant=group[0]
        for value in values.values():
            if isinstance(value,str) and (len(value)>32767 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',value)):
                raise ValueError('An original source field cannot be stored losslessly in Excel. Retain the original TSV and review this field before exporting.')
        original_row(sheet,[*[values.get(key) for key in keys],'; '.join(child.id for child in group),any(child.selected for child in group),variant.source_row,
                            Path(variant.source_file).name,variant.raw.get('_sheet',''),variant.assembly,analysis.id,variant.source_hash])
    style_table(sheet,name='SolideAnalysisRawData');sheet.sheet_state='hidden'


def _source_rows(workbook,variants,scoped,attachment_name,anchors,directory,analysis,*,norwegian=False):
    tr=report_text if norwegian else lambda value:value
    sheet=workbook.create_sheet('Kilder' if norwegian else 'Sources');title(sheet,tr('Source evidence'),7)
    sheet.merge_cells('A2:G2');sheet.cell(2,1,tr('Original rows: right-click a worksheet tab and choose Unhide → Raw data.')).alignment=Alignment(wrap_text=True)
    sheet.row_dimensions[2].height=30
    sheet.merge_cells('A3:G3');sheet.cell(3,1,tr('Source assertions are separate from laboratory assessments. Full responses and provenance are retained in hidden columns.')).alignment=Alignment(wrap_text=True)
    sheet.row_dimensions[3].height=30
    items=[]
    for variant in variants:
        for source,item in ordered_evidence(variant):items.append((variant,source,item,assess_evidence(variant,item,scoped)))
    records=[]
    for variant,source,item,assessment in items:
        url=_public_source_url(source,item.get('url',''));assertion=source_classification(item)
        records.append([variant.gene,variant.corrected_hgvs or variant.coding or variant.protein,source,
            tr(assessment.label),assertion[:120]+(('… (full assertion below)' if not norwegian else '… (se full tekst)') if len(assertion)>120 else ''),
            captured_time(item.get('captured_at','')),'',url,'',variant.id,item.get('accession',''),'',
            variant.source_row,Path(variant.source_file).name,analysis.id,variant.source_hash])
    row=_table_region(sheet,5,[tr(label) for label in ['Gene','Variant','Source','Match assessment','Source assertion','Captured (Oslo)','Evidence',
        'Source URL','Summary','Row ID','Accession','JSON response','Source row','Source file','Analysis ID','Source SHA256']],
        records,[13,28,17,22,32,21,18],'SolideAnalysisSources')
    for current,(variant,source,item,assessment) in enumerate(items,6):
        status_cell(sheet.cell(current,4),assessment.label);sheet.cell(current,4,tr(assessment.label));sheet.cell(current,6).number_format=EXCEL_DATETIME_FORMAT
        sheet.row_dimensions[current].height=_text_height(records[current-6][:7],[13,28,17,22,32,21,18],44)
        for column,field in [(9,_exported_response(source,item.get('summary',''))),
                             (12,json.dumps(_exported_response(source,item.get('raw',{})),ensure_ascii=False,default=str))]:
            if len(str(field or ''))<=32767:
                cell=sheet.cell(current,column,str(field or ''));cell.data_type='s';cell.quotePrefix=str(field or '').startswith(('=','+','-','@'))
        anchor=anchors.get((variant.id,source))
        if attachment_name and anchor:
            _link(sheet.cell(current,7),f"{attachment_name}#'{anchor[0].replace(chr(39),chr(39)*2)}'!A{anchor[1]}",tr('Open attachment'))
        elif _capture_export_allowed(variant,item,scoped):
            captures=_capture_records(item)
            capture=next((record['path'] for record in captures if valid_capture(record['path'])),None)
            if capture:_link(sheet.cell(current,7),_relative_target(capture,directory),tr('Open local capture'))
            elif _public_source_url(source,item.get('url','')):_link(sheet.cell(current,7),item['url'],tr('Open source'))
        if sheet.cell(current,7).value in (None,''):sheet.cell(current,7,tr('Review source') if assessment.label not in {'Not applicable','No match'} else tr(assessment.label))
    from openpyxl.utils import get_column_letter
    for column in range(8,17):sheet.column_dimensions[get_column_letter(column)].hidden=True
    if not records:
        sheet.merge_cells('A6:G6');sheet.cell(6,1,tr('No database searches have been recorded.'));sheet.row_dimensions[6].height=30
    metadata_row=max(row+2,9)
    for variant,source,item,assessment in items:
        full=json.dumps(_exported_response(source,item.get('raw',{})),ensure_ascii=False,default=str)
        summary=_exported_response(source,item.get('summary',''));assertion=source_classification(item)
        # Hidden literal rows have a stable source-result key; joining the
        # adjacent chunks reconstructs even multi-cell JSON exactly.
        for label,value in [(tr('JSON response'),full),(tr('Full summary'),summary),(tr('Full source assertion'),assertion)]:
            original_row(sheet,[None]*7+[variant.id,source,label])
            storage=sheet.max_row
            _literal_chunks(sheet,storage,11,value);sheet.row_dimensions[storage].hidden=True
        metadata_row=sheet.max_row+2
        if len(assertion)>120 or len(summary)>32767:
            metadata_row=evidence_field(sheet,metadata_row,f'{variant.gene} / {source}',variant.corrected_hgvs or variant.coding or variant.protein)
            if len(assertion)>120:metadata_row=evidence_field(sheet,metadata_row,tr('Full source assertion'),assertion)
            if len(summary)>32767:metadata_row=evidence_field(sheet,metadata_row,tr('Full summary'),summary)
    metadata_row=evidence_field(sheet,metadata_row,tr('Analysis'),f'{analysis.profile} / {analysis.assay}')
    metadata_row=evidence_field(sheet,metadata_row,tr('Source file'),Path(analysis.source_file or variants[0].source_file).name)
    if analysis.metadata:
        original_row(sheet,[None]*7+[analysis.id,'analysis',tr('Source metadata')])
        storage=sheet.max_row
        _literal_chunks(sheet,storage,11,json.dumps(analysis.metadata,ensure_ascii=False,default=str))
        sheet.row_dimensions[storage].hidden=True
    for variant in variants:
        group=annotation_group(variant)
        if not group:continue
        original_row(sheet,[None]*7+[variant.id,'annotation','Annotasjonsvalg'])
        storage=sheet.max_row
        mapping={**group,'gene_choice':variant.gene,'transcript_choice':variant.transcript,
                 'coding_choice':variant.coding,'protein_choice':variant.protein,'selected':variant.selected,
                 'resolution':variant.raw.get('_annotation_resolution',{})}
        _literal_chunks(sheet,storage,11,json.dumps(mapping,ensure_ascii=False))
        sheet.row_dimensions[storage].hidden=True
    sheet.print_area=f'A1:G{max(sheet.max_row,metadata_row)}'


def _capture_records(item):
    raw=item.get('raw',{});records=raw.get('screenshots',[]) or ([{'path':raw['screenshot']}] if raw.get('screenshot') else [])
    result=[];seen=set()
    for record in records:
        if not isinstance(record,dict):continue
        value=record.get('path')
        if not isinstance(value,str) or not value:continue
        key=str(Path(value).resolve())
        if key not in seen:result.append({'path':value,'label':record.get('label') or 'Evidence capture'});seen.add(key)
    return result


def _variant_attachment_name(variant,counts,used):
    gene=variant.gene or 'Variant'
    detail=(variant.protein or variant.corrected_hgvs or variant.coding) if counts[gene.casefold()]>1 else ''
    base=re.sub(r'[\\/*?:\[\]\x00-\x1f]','_',f'{gene} {detail}'.strip()) or 'Variant'
    name=base[:31];index=2
    while name.casefold() in used:
        suffix=f'-{index}';name=base[:31-len(suffix)]+suffix;index+=1
    used.add(name.casefold());return name


def _attachment_service_fields(source,raw):
    fields=service_fields(source,raw)
    if source!='BRCA Exchange':return fields
    release=raw.get('release') or {}
    if not isinstance(release,dict):release={}
    parts=[]
    for key,label in [('id','Release'),('version','Version'),('date','Date'),('archive','Archive'),('file','File')]:
        value=release.get(key)
        if value not in (None,''):parts.append(f'{label}: {value}')
    return [(label,'; '.join(parts) or 'See source response') if label=='Dataset release' else (label,value)
            for label,value in fields]


def _evidence_attachment(analysis,variants,scoped,directory):
    from collections import Counter
    eligible=[variant for variant in variants if any(evidence_is_current(variant,item,scoped.tissue_for(variant),scoped)
              for item in variant.evidence.values())]
    if not eligible:return None,{}
    workbook=Workbook();_workbook_marker(workbook);index=workbook.active;index.title='Index'
    title(index,f'{analysis.profile} evidence attachment',5)
    add_row(index,['Sample / dataset',analysis.patient,'Assay',analysis.assay])
    add_row(index,['Genome',analysis.assembly])
    records=[[variant.gene,variant.corrected_hgvs or variant.coding or variant.protein,'Open evidence'] for variant in eligible]
    _table_region(index,5,['Gene','Variant','Evidence'],records,[16,52,22],'SolideAttachmentIndex')
    anchors={};full_reports={};full_links=[];counts=Counter((variant.gene or 'Variant').casefold() for variant in eligible)
    used={'index','mtbp'}
    for index_row,variant in enumerate(eligible,6):
        name=_variant_attachment_name(variant,counts,used);sheet=workbook.create_sheet(name)
        title(sheet,f'{variant.gene}  {variant.corrected_hgvs or variant.coding or variant.protein}',8)
        _base_compact(sheet,[18]*8);internal_link(sheet.cell(2,1),'Index',label='Evidence index')
        internal_link(index.cell(index_row,3),name)
        row=4
        for label,value in [('Transcript',variant.transcript),('Original cDNA',variant.coding),('Protein',variant.protein),
                            ('Reviewed HGVS',variant.corrected_hgvs),('Reviewed allele',variant.controlled_genomic),
                            ('Genomic identity',' | '.join(filter(None,[variant.assembly,variant.locus,variant.ref,variant.alt]))),
                            ('Lab classification',variant.classification),('Decision',variant.report_decision),
                            ('Reviewer',variant.reviewer),('Reviewed at (Oslo)',captured_time(variant.reviewed_at)),
                            ('Assessment',variant.comment),('Source row',variant.source_row),('Source SHA256',variant.source_hash)]:
            if value not in (None,''):row=evidence_field(sheet,row,label,value)
        row+=1
        for source,item in ordered_evidence(variant):
            current=evidence_is_current(variant,item,scoped.tissue_for(variant),scoped)
            anchors[variant.id,source]=(name,row);row=section_heading(sheet,row,source)
            assessment=assess_evidence(variant,item,scoped)
            row=evidence_field(sheet,row,'Match assessment',assessment.label);status_cell(sheet.cell(row-1,2),assessment.label)
            for label,value in [('Source assertion',source_classification(item)),('Captured (Oslo)',captured_time(item.get('captured_at',''))),
                                ('Accession',item.get('accession','')),('Source URL',_public_source_url(source,item.get('url',''))),
                                ('Summary',_exported_response(source,item.get('summary','')))]:
                if value not in (None,''):row=evidence_field(sheet,row,label,value)
            for label,value in _attachment_service_fields(source,item.get('raw',{})):row=evidence_field(sheet,row,label,value)
            if not current or not _capture_export_allowed(variant,item,scoped):
                row=evidence_field(sheet,row,'Capture status','Outdated evidence: rerun before reporting.' if not current else
                                   'Capture excluded until the source result is verified or recovered.');row+=2;continue
            for record in _capture_records(item):
                if valid_capture(record['path']):row=put_image(sheet,record['path'],row,record['label'])
                else:row=evidence_field(sheet,row,'Capture status','Screenshot missing or unreadable: capture again.')
            full=item.get('raw',{}).get('patient_report_screenshot')
            if source=='MTBP' and isinstance(full,str) and full:
                key=str(Path(full).resolve());full_reports.setdefault(key,(full,item));full_links.append((sheet,row,key))
                sheet.cell(row,1,'Open full MTBP report');row+=2
            row+=2
        sheet.freeze_panes='A4';sheet.print_area=f'A1:H{max(sheet.max_row,row-1,max(sheet.row_dimensions,default=1))}'
    if full_reports:
        mtbp=workbook.create_sheet('MTBP');title(mtbp,'MTBP combined report',8);_base_compact(mtbp,[18]*8)
        internal_link(mtbp.cell(2,1),'Index',label='Evidence index');row=4;full_anchors={}
        for number,(key,(path,item)) in enumerate(full_reports.items(),1):
            full_anchors[key]=row;row=section_heading(mtbp,row,f'MTBP full report {number}')
            row=evidence_field(mtbp,row,'Captured (Oslo)',captured_time(item.get('captured_at','')))
            _link(mtbp.cell(row,1),_relative_target(path,directory),'Open original PNG');row+=2
            if valid_capture(path):row=put_image(mtbp,path,row,'MTBP report',crop_top=mtbp_content_top(path))
            else:row=evidence_field(mtbp,row,'Capture status','Full report missing or unreadable: capture again.')
            row+=2
        for sheet,link_row,key in full_links:internal_link(sheet.cell(link_row,1),'MTBP',full_anchors[key],label='Open full MTBP report')
        mtbp.print_area=f'A1:H{max(mtbp.max_row,row-1,max(mtbp.row_dimensions,default=1))}'
    index.print_area=f'A1:E{index.max_row}';return workbook,anchors


def _normalize_excel_views(workbook):
    """Repeated freeze_panes assignments must not retain obsolete selections."""
    from openpyxl.worksheet.views import Selection
    from openpyxl.utils.cell import coordinate_to_tuple
    for sheet in workbook:
        view=sheet.sheet_view;pane=view.pane
        if not pane:
            view.selection=[Selection(activeCell='A1',sqref='A1')];continue
        top,left=coordinate_to_tuple(pane.topLeftCell or 'A1')
        if pane.xSplit and pane.ySplit:
            from openpyxl.utils import get_column_letter
            cells=[('topRight',f'{get_column_letter(left)}1'),('bottomLeft',f'A{top}'),
                   ('bottomRight',pane.topLeftCell)]
            pane.activePane='bottomRight'
        else:
            pane.activePane='topRight' if pane.xSplit else 'bottomLeft'
            cells=[(pane.activePane,pane.topLeftCell)]
        view.selection=[Selection(pane=name,activeCell=cell,sqref=cell) for name,cell in cells]


def _save_report_bundle(main,attachment,report_path,attachment_path):
    """Stage both files completely, then publish the main report last."""
    staged=[];published_attachment=False
    try:
        for workbook,path in [(attachment,attachment_path),(main,report_path)]:
            if workbook is None:continue
            fd,name=tempfile.mkstemp(dir=path.parent,prefix='.solide-',suffix='.xlsx');os.close(fd)
            temporary=Path(name);staged.append((temporary,path));_normalize_excel_views(workbook);workbook.save(temporary)
        for temporary,path in staged:
            if path.exists():raise FileExistsError('An export version already exists. Retry export to create a new version.')
            os.replace(temporary,path)
            if attachment_path and path==attachment_path:published_attachment=True
    except Exception as error:
        if published_attachment and attachment_path and attachment_path.exists():attachment_path.unlink()
        if isinstance(error,PermissionError):
            raise PermissionError('Excel export could not save the report. Close any open report or check folder access, then retry export without rerunning database searches.') from error
        raise
    finally:
        for temporary,_ in staged:
            if temporary.exists():temporary.unlink()
        main.close()
        if attachment:attachment.close()


def export_analysis_compact(session:Session,analysis_id:str,directory:Path,include_evidence:bool=True)->ReportBundle:
    """Export exactly one analysis and its app assessments; never join by sample."""
    from .analyses import analysis_scoped_session
    scoped=analysis_scoped_session(session,analysis_id);analysis=scoped.analysis(analysis_id)
    if analysis.profile=='Unassigned':raise ValueError('Choose OPA or Ion Torrent for this analysis before exporting.')
    variants=scoped.analysis_variants(analysis_id)
    for variant in variants:
        if variant.selected:require_resolved_annotations(variant)
    directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    stem=re.sub(r'[^\w.-]','_',analysis.patient).strip('.')[:60] or 'source'
    profile=re.sub(r'[^\w.-]','_',analysis.profile)
    generated=oslo_now().replace(tzinfo=None);token=uuid.uuid4().hex[:10]
    base=f'{stem}_{profile}_{generated:%Y%m%d_%H%M%S}_{token}'
    report_path=directory/f'{base}_Report.xlsx';attachment_path=directory/f'{base}_Evidence.xlsx'
    attachment,anchors=_evidence_attachment(analysis,variants,scoped,directory) if include_evidence else (None,{})
    if attachment is None:attachment_path=None
    main=Workbook();_workbook_marker(main)
    _main_findings(main,analysis,variants,generated)
    _quality_sections(main,analysis,variants)
    _source_rows(main,variants,scoped,attachment_path.name if attachment_path else None,anchors,directory,analysis)
    _raw_analysis_sheet(main,variants,analysis)
    _save_report_bundle(main,attachment,report_path,attachment_path)
    return ReportBundle(report_path,attachment_path)


def _vpm_compact(variant,source,scoped):
    """Show source content; operational matching states remain in the app."""
    item=variant.evidence.get(source)
    if not item:return ''
    assessment=assess_evidence(variant,item,scoped).label
    if assessment=='Not applicable':
        return 'Ikke kjørt' if source=='SpliceAI' else ''
    if assessment=='No match':return 'Ingen treff'
    if assessment=='Outdated':return 'Utdatert'
    raw=item.get('raw',{})
    if source=='Mutalyzer' and evidence_is_current(variant,item,scoped.tissue_for(variant),scoped):
        checked=raw.get('nomenclature_assessment') or {}
        if checked.get('identity_confirmed') is True and checked.get('status') in {'Unchanged','Normalized'}:
            return 'Normalisert' if _vpm_normalized_name(variant,item) else 'Uendret'
        return 'Kontroller variant'
    if assessment not in {'Verified match','Review match'}:
        return 'Bilde mangler' if assessment=='Missing capture' else 'Kontroller resultat'
    assertion=' '.join(source_classification(item).split())
    if source=='COSMIC':
        identifiers=re.findall(r'COS[MV]\d+',str(item.get('accession','')))
        for key in ('mutation_id','cosmic_id','genomic_mutation_id','legacy_mutation_id'):
            identifiers.extend(re.findall(r'COS[MV]\d+',str(raw.get(key,''))))
        return ', '.join(dict.fromkeys(identifiers)) or 'Treff'
    if source=='SpliceAI':
        scores=raw.get('response',{}).get('scores') or []
        deltas=[float(row[key]) for row in scores for key in ('DS_AG','DS_AL','DS_DG','DS_DL') if row.get(key) is not None]
        return f'Maks. delta {max(deltas):.3f}' if deltas else 'Treff'
    if not assertion:
        for key in ('functional_relevance','oncogenicity','classification','expert_classification'):
            value=raw.get(key)
            if isinstance(value,str) and value.strip():assertion=' '.join(value.split());break
    if not assertion:
        match=re.search(r'(?:classification|oncogenicity|functional_relevance)=([^;|]+)',item.get('summary',''),re.I)
        if match:assertion=match[1].strip()
    if source=='MTBP':
        evidence=re.search(r'Evidence\s+[A-Z](?:\s*\([^)]*\))?',
                           str(raw.get('functional_evidence',''))+' '+item.get('summary',''),re.I)
        if evidence and evidence[0].casefold() not in assertion.casefold():
            assertion=' – '.join(filter(None,[assertion,' '.join(evidence[0].split())]))
    return assertion or 'Treff'


def _vpm_source_url(variant,source,scoped):
    item=variant.evidence.get(source)
    if not item or source=='MTBP':return ''
    if assess_evidence(variant,item,scoped).label not in {'Verified match','Review match'}:return ''
    return _public_source_url(source,item.get('url',''))


def _vpm_identity(variant):
    return display_identity(variant)


def _vpm_normalized_name(variant,item):
    raw=item.get('raw',{});checked=raw.get('nomenclature_assessment',{})
    if checked.get('identity_confirmed') is not True:return False
    if checked.get('status')=='Normalized':return True
    original=variant.coding if ':' in variant.coding else f'{variant.transcript}:{variant.coding}'
    return (checked.get('status')=='Unchanged' and variant.nomenclature_verified
            and bool(variant.corrected_hgvs) and variant.corrected_hgvs!=original
            and checked.get('suggested_hgvs')==raw.get('original')==variant.corrected_hgvs)


def _vpm_source_visible(variant,source,scoped):
    item=variant.evidence.get(source,{})
    if not item or item.get('status')=='not_applicable':return False
    if source=='Mutalyzer':
        return (evidence_is_current(variant,item,scoped.tissue_for(variant),scoped)
                and _vpm_normalized_name(variant,item)
                and _capture_export_allowed(variant,item,scoped))
    if source=='SpliceAI':return assess_evidence(variant,item,scoped).label=='Verified match'
    if source=='BRCA Exchange':
        return variant.gene in {'BRCA1','BRCA2'} and assess_evidence(variant,item,scoped).label=='Verified match'
    return True


def _vpm_variant_name(variant,index,counts,used):
    gene=variant.gene or f'Variant {index}'
    detail=''
    if counts[(variant.gene or 'Variant').casefold()]>1:
        detail=(variant.protein or variant.coding).split(':',1)[-1]
    base=re.sub(r'[\\/*?:\[\]\x00-\x1f]+','_',f'{gene} {detail}'.strip()) or f'Variant {index}'
    name=base[:31];suffix=2
    while name.casefold() in used:
        tail=f'-{suffix}';name=base[:31-len(tail)]+tail;suffix+=1
    used.add(name.casefold());return name


def _vpm_overview(workbook,analysis,selected,scoped,names,generated):
    import textwrap
    from openpyxl.utils import get_column_letter
    sheet=workbook.active;sheet.title='Oversikt';vpm_base_sheet(sheet,13)
    sheet.sheet_properties.tabColor=VPM_COLORS['navy']
    vpm_title(sheet,f'Solide – {analysis.patient}',13)
    vpm_info_row(sheet,3,'Prøve',analysis.patient,end_column=3)
    vpm_info_row(sheet,4,'Opprettet (Oslo)',generated,end_column=3);sheet.cell(4,3).number_format=EXCEL_DATETIME_FORMAT
    sheet.merge_cells('D3:M4')
    sheet.cell(3,4).alignment=Alignment(vertical='top',wrap_text=True)
    for row in sheet.iter_rows(min_row=3,max_row=4,min_col=4,max_col=13):
        for cell in row:cell.fill=PatternFill('solid',fgColor=VPM_COLORS['pale_orange'])
    for row in (3,4):sheet.row_dimensions[row].height=24
    vpm_section(sheet,6,'Varianter og kildefunn',13)
    for column,label in enumerate(['Gen','HGVSc','HGVSp','Kildefunn','Kommentar',*SOURCE_ORDER],1):
        vpm_header(sheet.cell(7,column),label)
    widths=[14,31,25,52,32,*([22]*len(SOURCE_ORDER))]
    for row,variant in enumerate(selected,8):
        labels=[_vpm_compact(variant,source,scoped) for source in SOURCE_ORDER]
        coding,protein,_=_vpm_identity(variant)
        values=[variant.gene,coding,protein,
                '\n'.join(f'{source}: {label}' for source,label in zip(SOURCE_ORDER,labels) if label),
                '',*labels]
        for column,value in enumerate(values,1):
            cell=sheet.cell(row,column,safe_text(value));cell.border=vpm_border()
            cell.alignment=Alignment(vertical='top',wrap_text=True)
            cell.fill=PatternFill('solid',fgColor=VPM_COLORS['pale_blue'] if (row-8)%2==0 else VPM_COLORS['white'])
        for column,source in enumerate(SOURCE_ORDER,6):
            url=_vpm_source_url(variant,source,scoped)
            if url:
                sheet.cell(row,column).hyperlink=url;sheet.cell(row,column).style='Hyperlink'
                sheet.cell(row,column).alignment=Alignment(vertical='top',wrap_text=True)
        # Archer reapplies the alternating fill after hyperlink styling.
        for cell in sheet[row][:13]:
            cell.fill=PatternFill('solid',fgColor=VPM_COLORS['pale_blue'] if (row-8)%2==0 else VPM_COLORS['white'])
            cell.border=vpm_border()
        lines=max(sum(max(1,len(textwrap.wrap(line,max(1,width-7)))) for line in str(value or '').split('\n'))
                  for value,width in zip(values,widths))
        sheet.row_dimensions[row].height=min(409,max(72,lines*16+12))
    for column,width in enumerate(widths,1):sheet.column_dimensions[get_column_letter(column)].width=width
    sheet.auto_filter.ref=f'A7:M{max(7,7+len(selected))}';sheet.freeze_panes='A8'
    sheet.print_area=f'A1:M{max(13,8+len(selected))}'
    if not selected:
        sheet.merge_cells('A8:M8');sheet.cell(8,1,'Ingen varianter valgt.')
        sheet.cell(8,1).alignment=Alignment(wrap_text=True);sheet.row_dimensions[8].height=42


def _vpm_variant_sheet(workbook,name,variant,scoped,include_evidence,full_reports):
    sheet=workbook.create_sheet(name);vpm_base_sheet(sheet)
    coding,protein,transcript=_vpm_identity(variant)
    vpm_title(sheet,' | '.join(filter(None,[variant.gene,coding,protein])))
    sheet.merge_cells('A3:L3')
    metadata=[protein,'Transkript: '+transcript if transcript else '',
              ' | '.join(filter(None,[variant.assembly,variant.locus,variant.ref,variant.alt])),
              f'AF {variant.af_percent:.2f}%' if variant.af_percent is not None else '',
              f'Dekning {variant.coverage:g}' if variant.coverage is not None else '',variant.kind]
    sheet.cell(3,1,' | '.join(filter(None,metadata)))
    sheet.cell(3,1).font=Font(color=VPM_COLORS['muted'],italic=True)
    sheet.cell(3,1).alignment=Alignment(wrap_text=True)
    sheet.row_dimensions[3].height=_text_height([sheet.cell(3,1).value],[145],30)
    row=5;anchors={}
    for source in SOURCE_ORDER:
        item=variant.evidence.get(source,{})
        if not _vpm_source_visible(variant,source,scoped):continue
        row=vpm_section(sheet,row,source,source=source)
        label=_vpm_compact(variant,source,scoped)
        if label:sheet.cell(row-1,1,source+' – '+label)
        url=_vpm_source_url(variant,source,scoped)
        if url:sheet.cell(row-1,1).hyperlink=url
        anchors[variant.id,source]=(name,row-1)
        allowed=bool(item) and _capture_export_allowed(variant,item,scoped)
        records=[record for record in _capture_records(item) if valid_capture(record['path'])] if allowed else []
        full=item.get('raw',{}).get('patient_report_screenshot')
        has_full=source=='MTBP' and allowed and isinstance(full,str) and valid_capture(full)
        if has_full:
            full_key=str(Path(full).resolve())
            records=[record for record in records if str(Path(record['path']).resolve())!=full_key]
        if not include_evidence:message='Bilder er utelatt.'
        elif not allowed and item:message=_vpm_compact(variant,source,scoped)+': bilde utelatt.'
        elif not records and has_full:message='Se MTBP-arket.'
        elif not records:message='Bilde mangler.'
        else:message=''
        if message:
            sheet.merge_cells(start_row=row,start_column=1,end_row=row,end_column=12)
            sheet.cell(row,1,message).alignment=Alignment(wrap_text=True,vertical='top')
            sheet.cell(row,1).fill=PatternFill('solid',fgColor=VPM_COLORS['pale_orange'])
            sheet.row_dimensions[row].height=42;row+=2
        else:
            for record in records:
                sheet.merge_cells(start_row=row,start_column=1,end_row=row,end_column=12)
                caption=sheet.cell(row,1,safe_text(capture_caption(record['label'])))
                caption.font=Font(bold=True,color=VPM_COLORS['navy'])
                caption.fill=PatternFill('solid',fgColor=VPM_COLORS['pale_blue'])
                row=vpm_put_image(sheet,record['path'],row+1)+1
        if include_evidence and has_full:
            full_reports.setdefault(str(Path(full).resolve()),full)
            row+=1
    sheet.freeze_panes='A5';sheet.print_area=f'A1:L{max(5,row-1)}'
    return anchors


def _vpm_mtbp_sheet(workbook,analysis,full_reports):
    if not full_reports:return
    sheet=workbook.create_sheet('MTBP',1);vpm_base_sheet(sheet)
    sheet.sheet_view.showGridLines=True;sheet.sheet_properties.tabColor=VPM_COLORS['muted']
    sheet.cell(1,1,safe_text(analysis.patient)).font=Font(size=18,bold=True,color=VPM_COLORS['navy'])
    sheet.column_dimensions['A'].width=max(18,len(analysis.patient)+4)
    for row in range(2,6):
        for column in range(1,13):sheet.cell(row,column).fill=PatternFill('solid',fgColor=VPM_COLORS['pale_blue'])
    row=6
    for path in full_reports.values():
        row=vpm_section(sheet,row,'MTBP – samlet rapport')
        row=vpm_put_image(sheet,path,row,crop_top=mtbp_content_top(path))+1
    sheet.print_area=f'A1:L{row}'


def export_analysis(session:Session,analysis_id:str,directory:Path,include_evidence:bool=True)->ReportBundle:
    """Export one analysis in the current Archer/VPM single-workbook layout.

    Only the app's selected rows appear in Overview and variant tabs. All
    supplied rows and full responses remain in hidden lossless source storage.
    """
    from collections import Counter
    from .analyses import analysis_scoped_session
    scoped=analysis_scoped_session(session,analysis_id);analysis=scoped.analysis(analysis_id)
    if analysis.profile=='Unassigned':raise ValueError('Choose OPA or Ion Torrent for this analysis before exporting.')
    variants=scoped.analysis_variants(analysis_id);selected=[variant for variant in variants if variant.selected]
    for variant in selected:require_resolved_annotations(variant)
    directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    generated=oslo_now().replace(tzinfo=None)
    stem=re.sub(r'[^\w.-]','_',analysis.patient).strip('.')[:60] or 'source'
    profile=re.sub(r'[^\w.-]','_',analysis.profile)
    path=directory/f'{stem}_{profile}_{generated:%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:10]}_Report.xlsx'
    workbook=Workbook();_workbook_marker(workbook);workbook.properties.version='6'
    used={'oversikt','mtbp','kilder','rådata','overview','sources','raw data','quality'}
    counts=Counter((variant.gene or 'Variant').casefold() for variant in selected)
    names={variant.id:_vpm_variant_name(variant,index,counts,used) for index,variant in enumerate(selected,1)}
    full_reports={};anchors={}
    try:
        _vpm_overview(workbook,analysis,selected,scoped,names,generated)
        for variant in selected:
            anchors.update(_vpm_variant_sheet(workbook,names[variant.id],variant,scoped,include_evidence,full_reports))
        _vpm_mtbp_sheet(workbook,analysis,full_reports)
        _source_rows(workbook,variants,scoped,None,{},directory,analysis,norwegian=True)
        sources=workbook['Kilder']
        sources.sheet_state='hidden'
        _raw_analysis_sheet(workbook,variants,analysis,norwegian=True)
        # Hidden provenance is literal storage, not an Excel structured table.
        for sheet in (sources,workbook['Rådata']):
            sheet.tables.clear();sheet.auto_filter.ref=None
            for row in sheet:
                for cell in row:
                    if cell.hyperlink:cell.hyperlink=None
        _save_report_bundle(workbook,None,path,None)
    except Exception:
        workbook.close();raise
    return ReportBundle(path)
