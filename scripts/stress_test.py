"""Opt-in synthetic stress run. Never uses real patients, accounts or live databases.

Run: python scripts/stress_test.py [--phase import|gui|all] [--actions 1200]
Fixtures, timings and findings are retained under artifacts/stress/<Oslo timestamp>.
"""
import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import argparse,csv,hashlib,io,json,random,sys,time,traceback
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
from solide.importing import load_file
from solide.localtime import oslo_now

parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('--phase',choices=['import','gui','all'],default='all')
parser.add_argument('--actions',type=int,default=1200)
args=parser.parse_args()
root=Path(__file__).resolve().parents[1]/'artifacts'/'stress'/oslo_now().strftime('%Y%m%d-%H%M%S')
root.mkdir(parents=True)
(root.parent/'latest.txt').write_text(str(root),encoding='utf8')
inputs=root/'inputs';inputs.mkdir()
rng=random.Random(20261018)
checks=[]
def check(name,fn):
    started=time.monotonic()
    try:
        details=fn();record={'name':name,'passed':True,'details':details}
    except Exception as exc:
        record={'name':name,'passed':False,'error':str(exc),'traceback':traceback.format_exc()}
    record['seconds']=round(time.monotonic()-started,3);checks.append(record)
    (root/'results.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps({k:v for k,v in record.items() if k!='traceback'},ensure_ascii=False),flush=True)

def tsv(header,rows,meta='##reference=hg19\n##sampleNames=00OUM00001\n'):
    stream=io.StringIO(newline='');stream.write(meta)
    writer=csv.writer(stream,delimiter='\t',lineterminator='\n');writer.writerow(header);writer.writerows(rows)
    return stream.getvalue()
ion=['Type','Genes','Transcript','Coding','Coverage','Allele Frequency','Genotype','Comment']
opa=['Type','Gene','Gene Isoform','Nucleotide Change','Coverage','Allele Frequency %','Call','Comment']
base=['SNV','EGFR','NM_005228.5','c.2573T>G','2000','0.25','PRESENT','SYNTHETIC ONLY']
fixtures=[]
def fixture(name,payload,expected='accept',verify=None):
    path=inputs/(name+'.tsv');path.write_bytes(payload if isinstance(payload,bytes) else payload.encode('utf8'))
    fixtures.append((path,expected,verify));return path
valid=[]
for i in range(24):
    header=ion if i%2 else opa;rows=[]
    for j in range(12):
        row=base.copy();row[1]=['EGFR','TP53','BRCA1','PTEN','MET','MTBP','Oversikt','Rådata','A/B:C?D'][j%9]
        row[3]=['c.2573T>G','c.743G>A','c.68_69del','c.254-1G>A'][j%4]
        row[4]=['0','499','500','-1','NaN','inf','2000'][j%7]
        row[5]=['0.1','0.25','-1','101','NaN','1e309','12,5'][j%7]
        row[6]='ABSENT' if j%5==0 else 'PRESENT'
        row[7]=['SYNTHETIC ONLY','=1+1','+SUM(A1:A2)','@example','Æ Ø Å – tab\there'][j%5]
        rows.append(row)
    valid.append(fixture(f'valid-{i:02}',tsv(header,rows,meta=f'##reference=hg19\n##sampleNames=00OUM{i+1:05}\n')))
fixture('empty','',expected='reject')
fixture('wrong-header','Type\tSomething\nSNV\tEGFR\n',expected='reject')
fixture('duplicate-columns',tsv(ion+['Genes'],[base+['TP53']]),expected='reject')
fixture('two-metadata-samples',tsv(ion,[base],meta='##sampleNames=00OUM00001,00OUM00002\n'),expected='reject')
fixture('two-row-samples',tsv(ion+['Sample ID'],[base+['00OUM00001'],base+['00OUM00002']]),expected='reject')
fixture('invalid-encoding',b'\x81\xff',expected='reject')
fixture('too-many-cells',tsv(ion,[base+['UNNAMED DATA MUST NOT DISAPPEAR']]),expected='reject')
fixture('two-concatenated-tables',tsv(ion,[base])+tsv(ion,[base],meta='##sampleNames=00OUM00002\n'),expected='reject')
fixture('unclosed-quote',tsv(ion,[])+ '\t'.join(base[:-1])+'\t"unfinished\n',expected='reject')
multiline='First line\nSecond line\twith tab and "quote"'
fixture('quoted-multiline',tsv(ion,[base[:-1]+[multiline]]),verify=lambda result: result.variants[0].raw['Comment']==multiline)
fixture('cp1252',tsv(ion,[base[:-1]+['Ærlig prøve']]).encode('cp1252'))
fixture('bom',b'\xef\xbb\xbf'+tsv(ion,[base]).encode('utf8'))
fixture('missing-oum',tsv(opa,[base],meta='##reference=hg19\n'))
fixture('mixed-case-oum',tsv(opa,[base],meta='##sampleNames=00oum00099\n'))
fixture('extreme-intron-offset',tsv(ion,[base[:3]+['c.100+'+'9'*5000+'A>G']+base[4:]]))
large=fixture('large-5000',tsv(ion,[base[:-1]+[f'SYNTHETIC row {i}'] for i in range(5000)]))
hashes={str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p,_,_ in fixtures}

def parse_one(path,expected,verify):
    try:result=load_file(path)
    except (ValueError,UnicodeError,csv.Error) as exc:
        if expected=='reject':return {'rejected':type(exc).__name__}
        raise
    assert expected!='reject','Damaged/ambiguous file was accepted'
    if verify:assert verify(result),'Imported content differs from original cell'
    from solide.quality import review_reasons
    for variant in result.variants:review_reasons(variant)
    return {'rows':len(result.variants),'profile':result.profile}
if args.phase in {'import','all'}:
    for path,expected,verify in fixtures:check('import/'+path.stem,lambda p=path,e=expected,v=verify:parse_one(p,e,v))
    def fuzz():
        accepted=rejected=0
        for i in range(300):
            header=ion.copy();rows=[]
            for j in range(rng.randrange(1,12)):
                row=base.copy();row[7]=''.join(rng.choice('abcÆØÅ\t\n" =+-@') for _ in range(rng.randrange(1,120)))
                rows.append(row)
            path=inputs/f'fuzz-{i:03}.tsv';path.write_text(tsv(header,rows),encoding='utf8',newline='')
            result=load_file(path);accepted+=1
            assert len(result.variants)==len(rows),(i,len(rows),len(result.variants))
            assert [v.raw['Comment'] for v in result.variants]==[row[7] for row in rows],f'Cell loss in generated case {i}'
        return {'accepted':accepted,'rejected':rejected}
    check('import/300-quoted-cell-roundtrips',fuzz)

if args.phase in {'gui','all'}:
    from PyQt6.QtWidgets import QApplication,QFileDialog,QMessageBox,QPushButton
    from PyQt6.QtCore import Qt
    from PyQt6.QtTest import QTest
    from solide.gui import MainWindow
    from solide.reporting import export_analysis
    from openpyxl import load_workbook
    app=QApplication.instance() or QApplication([])
    errors=[];qt_errors=[]
    old_hook=sys.excepthook
    sys.excepthook=lambda kind,value,tb:qt_errors.append(''.join(traceback.format_exception(kind,value,tb)))
    with patch('solide.gui.Path.home',return_value=root/'isolated-home'),patch('solide.gui.read_password',return_value=''):
        view=MainWindow()
    view.config_path=root/'isolated-settings.json';view.show();app.processEvents()
    warnings=patch.object(QMessageBox,'warning',side_effect=lambda parent,title,message:errors.append(str(message)))
    warnings.start()
    def damaged_settings():
        folder=root/'damaged-settings';folder.mkdir()
        settings=folder/'.solide'/'ui_config.json';settings.parent.mkdir()
        payloads=[[],None,42,{'sources':None},{'background':'false'},{'accounts':[]},
            {'accounts':{'Franklin':False}},{'login_checks':[]},{'login_checks':{'Franklin':[]}},
            {'login_checks':{'Franklin':{'state':42,'time':[]}}}]
        for i,payload in enumerate(payloads):
            text=json.dumps(payload);settings.write_text(text,encoding='utf8')
            (folder/f'case-{i:02}.json').write_text(text,encoding='utf8')
            with patch('solide.gui.Path.home',return_value=folder),patch('solide.gui.read_password',return_value=''):
                candidate=MainWindow()
                try:assert candidate.nav.count()==4
                finally:candidate.close()
            assert settings.read_text(encoding='utf8')==text
        assert not qt_errors,qt_errors[-1:]
        return {'malformed_settings':len(payloads),'started':len(payloads),'source_files_unchanged':True}
    check('gui/startup-with-damaged-settings',damaged_settings)
    def wait(limit=30):
        end=time.monotonic()+limit
        while view.active and time.monotonic()<end:app.processEvents();time.sleep(.002)
        assert not view.active,'Worker did not finish'
        app.processEvents()
    def import_batch():
        paths=[str(p) for p in valid]
        with patch.object(QFileDialog,'getOpenFileNames',return_value=(paths,'')):
            QTest.mouseClick(view.import_btn,Qt.MouseButton.LeftButton);wait()
        assert len(view.session.analyses)==24
        assert len(view.session.variants)==288
        with patch.object(QFileDialog,'getOpenFileNames',return_value=(paths,'')):
            view.import_files();wait()
        assert len(view.session.analyses)==24 and len(view.session.variants)==288
        return {'files':24,'rows':288,'duplicate_imports':24}
    check('gui/import-and-repeat-24-files',import_batch)
    def random_actions():
        buttons={b.text():b for b in view.findChildren(QPushButton)}
        operations=['tab','sample','filter','sort','select','check','font','details','rows','copy','view']
        for i in range(args.actions):
            action=rng.choice(operations)
            if action=='tab':view.nav.setCurrentRow(rng.randrange(4))
            elif action=='sample':
                index=view.sample_list.model().index(rng.randrange(view.sample_list.count()),0)
                view.sample_list.scrollTo(index);app.processEvents()
                QTest.mouseClick(view.sample_list.viewport(),Qt.MouseButton.LeftButton,pos=view.sample_list.visualRect(index).center())
            elif action=='filter':view.search.setText(rng.choice(['','TP53','no such variant','[.*]','ÆØÅ','c.68','00OUM']))
            elif action=='sort':view.proxy.sort(rng.randrange(13),rng.choice([Qt.SortOrder.AscendingOrder,Qt.SortOrder.DescendingOrder]))
            elif action=='select':
                name=rng.choice(['Select visible','Clear visible']);QTest.mouseClick(buttons[name],Qt.MouseButton.LeftButton)
            elif action=='check' and view.proxy.rowCount():
                index=view.proxy.index(rng.randrange(view.proxy.rowCount()),0)
                view.proxy.setData(index,rng.choice([Qt.CheckState.Checked,Qt.CheckState.Unchecked]),Qt.ItemDataRole.CheckStateRole)
            elif action=='font':view.text_size.setCurrentIndex(rng.randrange(4))
            elif action=='details':view.details_toggle.setChecked(not view.details_toggle.isChecked())
            elif action=='rows':view.show_all_rows.toggle()
            elif action=='view':view.workspace_view.setCurrentIndex(rng.randrange(2))
            elif action=='copy' and view.proxy.rowCount():
                view.variant_table.selectRow(rng.randrange(view.proxy.rowCount()))
                QTest.keyClick(view.variant_table,Qt.Key.Key_C,Qt.KeyboardModifier.ControlModifier)
            app.processEvents()
            active=view.session.analysis(view.active_analysis_id)
            assert all(v.analysis_id==active.id and v.patient==active.patient for v in view.variant_model.variants),'Cross-sample rows'
            assert len({v.id for v in view.session.variants})==288
            assert not qt_errors,qt_errors[-1:]
        view.grab().save(str(root/'random-actions.png'))
        return {'actions':args.actions,'uncaught_qt_errors':len(qt_errors),'files':len(view.session.analyses)}
    check('gui/random-clicks-and-keyboard',random_actions)
    def reports():
        view.search.clear()
        for i in range(6):
            view.analysis_selector.setCurrentIndex(i)
            view.select_visible(True)
            before=len(list(root.glob('*.xlsx')))
            with patch.object(QFileDialog,'getExistingDirectory',return_value=str(root)):
                view.export();wait()
            assert len(list(root.glob('*.xlsx')))==before+1,errors
        for path in root.glob('*.xlsx'):
            book=load_workbook(path)
            try:
                assert book['Rådata'].max_row==13
                assert book['Kilder'].sheet_state=='hidden'
                for sheet in book:
                    assert len({s.title.casefold() for s in book})==len(book.sheetnames)
                    for row in sheet:
                        assert all(cell.data_type!='f' for cell in row),'Source formula executed'
            finally:book.close()
        return {'reports':len(list(root.glob('*.xlsx'))),'literal_formula_cells':True}
    check('gui/export-six-analyses',reports)
    def busy_actions():
        from solide.gui import Worker
        view.search.clear();view.select_visible(True)
        original_ids=[a.id for a in view.session.analyses]
        original_active=view.active_analysis_id
        from threading import Event
        release=Event()
        worker=Worker(lambda:release.wait(3))
        try:
            view.launch(worker)
            for _ in range(30):
                QTest.mouseClick(view.clear_btn,Qt.MouseButton.LeftButton)
                QTest.mouseClick(view.run_btn,Qt.MouseButton.LeftButton)
                QTest.mouseClick(view.report_btn,Qt.MouseButton.LeftButton)
                view.nav.setCurrentRow(rng.randrange(4));app.processEvents()
            assert [a.id for a in view.session.analyses]==original_ids
            assert view.active_analysis_id==original_active
            assert not view.close(),'Closed while worker was running'
        finally:release.set();wait()
        assert view.import_btn.isEnabled() and view.clear_btn.isEnabled()
        return {'blocked_clicks':90,'close_during_work':'blocked','controls_recovered':True}
    check('gui/busy-double-clicks-and-close',busy_actions)
    def failed_export():
        target=root/'not-a-directory';target.write_text('SYNTHETIC export error',encoding='utf8')
        old_files=set(root.glob('*.xlsx'));old_errors=len(errors)
        with patch.object(QFileDialog,'getExistingDirectory',return_value=str(target)):
            view.export();wait()
        assert len(errors)==old_errors+1,'Filesystem error was not reported'
        assert set(root.glob('*.xlsx'))==old_files
        assert not list(root.glob('.solide-*.xlsx'))
        with patch.object(QFileDialog,'getExistingDirectory',return_value=str(root)):
            view.export();wait()
        assert len(set(root.glob('*.xlsx'))-old_files)==1,'Retry did not recover'
        return {'filesystem_error':'reported','partial_files':0,'retry':'succeeded'}
    check('gui/export-failure-and-retry',failed_export)
    def queue_failure():
        import solide.gui as gui
        view.search.clear();view.select_visible(True)
        for name,c in view.source_checks.items():c.setChecked(name=='ClinVar')
        selected=[v for v in view.active_session().variants if v.selected]
        def failing(session,sources,root,control,emit,progress,background,**kwargs):
            variant=next(v for v in session.variants if v.selected)
            control.wait(.03)
            emit(variant.id,'ClinVar',{'database':'ClinVar','status':'not_found','summary':'SYNTHETIC negative',
                'fingerprint':variant.fingerprint('Other'),'raw':{}})
            raise RuntimeError('SYNTHETIC database disconnect')
        with patch.object(gui,'run_queue',failing),patch.object(gui,'browser_credentials',return_value={}):
            view.start_queue('all');wait()
        assert view.run_state=='Failed'
        assert len(view.run_done)==1 and selected[0].evidence['ClinVar']['status']=='not_found'
        assert view.run_btn.isEnabled() and view.import_btn.isEnabled()
        def controlled(session,sources,root,control,emit,progress,background,**kwargs):
            for i,variant in enumerate(v for v in session.variants if v.selected):
                control.wait(.04)
                emit(variant.id,'ClinVar',{'database':'ClinVar','status':'timeout' if i==0 else 'not_found',
                    'summary':'SYNTHETIC timeout' if i==0 else 'SYNTHETIC no match',
                    'fingerprint':variant.fingerprint('Other'),'raw':{}})
        with patch.object(gui,'run_queue',controlled),patch.object(gui,'browser_credentials',return_value={}):
            view.start_queue('all')
            QTest.mouseClick(view.pause_btn,Qt.MouseButton.LeftButton);app.processEvents()
            assert view.control.paused.is_set()
            QTest.mouseClick(view.pause_btn,Qt.MouseButton.LeftButton);app.processEvents()
            assert not view.control.paused.is_set()
            end=time.monotonic()+3
            while len(view.run_done)<2 and time.monotonic()<end:app.processEvents();time.sleep(.005)
            view.stop_queue();wait()
            assert view.run_state=='Stopped'
            assert 2<=len(view.run_done)<len(view.run_plan)
            preserved=len(view.run_done)
            view.start_queue('pending');wait()
            assert view.run_state=='Completed'
        assert not qt_errors,qt_errors[-1:]
        return {'disconnect':'partial results retained','pause_resume':'passed','stop':'passed',
                'retained_on_stop':preserved,'new_run':'completed','timeouts':'reported'}
    check('gui/database-disconnect-pause-stop-restart',queue_failure)
    def large_import():
        view.clear_files();app.processEvents()
        start=time.monotonic()
        with patch.object(QFileDialog,'getOpenFileNames',return_value=([str(large)],'')):
            view.import_files();wait(90)
        assert len(view.session.variants)==5000
        elapsed=time.monotonic()-start
        view.search.setText('EGFR');app.processEvents()
        assert view.proxy.rowCount()==5000
        view.grab().save(str(root/'large-5000.png'))
        assert not qt_errors,qt_errors[-1:]
        return {'rows':5000,'import_seconds':round(elapsed,3)}
    check('gui/5000-row-file',large_import)
    temporary=view.work_root()
    view.close();app.processEvents();sys.excepthook=old_hook;warnings.stop()
    check('gui/cleanup',lambda: {'cleaned':not temporary.exists()} if not temporary.exists() else (_ for _ in ()).throw(AssertionError('Temporary files retained')))

def unchanged():
    assert all(hashlib.sha256(Path(p).read_bytes()).hexdigest()==digest for p,digest in hashes.items())
    return {'original_fixture_files':len(hashes),'unchanged':True}
check('source-files-unchanged',unchanged)
failed=[r['name'] for r in checks if not r['passed']]
(root/'README.md').write_text('# Syntetisk stresstest\n\nAlle filer er kunstige. OUM-numre med 00 og alle variantrader er testdata.\n'
    '\nIngen virkelige pasienter, kontoer eller eksterne databasesøk er brukt.\n'
    '\nSe results.json for kontroller, tider og eventuelle feil.\n',encoding='utf8')
print(json.dumps({'root':str(root),'checks':len(checks),'failed':failed}),flush=True)
raise SystemExit(bool(failed))
